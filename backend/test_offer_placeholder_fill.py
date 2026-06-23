"""Tests du moteur de remplissage « à variables » {{ }} (Offre_Technique_template).

Aucun appel LLM : on fournit des `sections` synthétiques et on vérifie que le .docx
produit est complet — plus aucun placeholder résiduel, nom du client substitué partout,
titre dans le header, listes éclatées, tableau Fonctionnalité→Composant inséré, planning
mis à jour (J/H par activité + TOTAL), tableau équipe dynamisé seulement si des CV sont
fournis.
"""
import os
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn

from modules.uc10_presales.offer_generator import (
    OfferGenerator, _all_text_roots, _find_team_table, _find_planning_table,
    _extract_planning_activities, _placeholders_present, _cover_date_parts,
)
from modules.uc10_presales.template_validator import validate_template, make_validation

PATH = "../data/ged/offres-techniques/TEMPLATE/Offre_Technique_template.docx"

pytestmark = pytest.mark.skipif(
    not os.path.exists(PATH), reason="template Offre_Technique_template.docx absent"
)


def _sections(n_planning: int = 30) -> dict:
    return {
        "titre_projet": "Conception d'une plateforme de gestion documentaire",
        "expression_besoins": ["Besoin 1.", "Besoin 2.", "Besoin 3."],
        "presentation_reponse": ["Proposition 1.", "Proposition 2.", "Proposition 3."],
        "fonctionnalites": ["Authentification", "GED", "Recherche", "Reporting", "Admin"],
        "repartition": [
            {"fonctionnalite": "Authentification", "composant": "Sécurité / Backend"},
            {"fonctionnalite": "GED", "composant": "Backend API"},
            {"fonctionnalite": "Recherche", "composant": "Moteur d'index"},
        ],
        # plus d'items que d'activités du modèle : J/H pris dans l'ordre, surplus ignoré
        "planning": [{"phase": "", "activite": "", "jh": "7"} for _ in range(n_planning)],
    }


def _fill(client="ACME Corp", cvs=None, abes=None) -> Document:
    gen = OfferGenerator(llm=None)
    doc = Document(PATH)
    data = gen._fill_placeholder_template(doc, None, _sections(), client, cvs or [], abes or [])
    return Document(BytesIO(data))


def _all_paragraph_texts(doc) -> list[str]:
    out = []
    for root in _all_text_roots(doc):
        for p_el in root.findall(".//" + qn("w:p")):
            out.append("".join((t.text or "") for t in p_el.findall(".//" + qn("w:t"))))
    return out


def test_template_is_placeholder_and_valid():
    doc = Document(PATH)
    assert _placeholders_present(doc), "le template doit être détecté « à variables »"
    rep = make_validation("(test)", PATH, validate_template(PATH))
    assert rep.ok, f"template invalide : {[c.label for c in rep.checks if not c.ok]}"


def test_no_leftover_placeholders():
    out = _fill()
    leftovers = [t for t in _all_paragraph_texts(out) if "{{" in t or "}}" in t]
    assert not leftovers, f"placeholders résiduels : {leftovers}"


def test_client_name_substituted_everywhere():
    out = _fill(client="ACME Corp")
    joined = " ".join(_all_paragraph_texts(out))
    assert "ACME Corp" in joined
    # le nom legacy de gabarit ne doit plus traîner
    assert "{{Nom client" not in joined


def test_title_in_header():
    out = _fill()
    titre = "plateforme de gestion documentaire"
    found = False
    for s in out.sections:
        el = s.header._element
        htxt = "".join((t.text or "") for t in el.findall(".//" + qn("w:t")))
        if titre in htxt:
            found = True
            break
    assert found, "le titre de l'offre doit figurer dans un header de page"


def test_lists_expanded():
    out = _fill()
    txts = [p.text.strip() for p in out.paragraphs]
    for b in ("Besoin 1.", "Besoin 2.", "Besoin 3."):
        assert b in txts, f"besoin manquant : {b}"
    for f in ("Authentification", "GED", "Recherche", "Reporting", "Admin"):
        assert f in txts, f"fonctionnalité manquante : {f}"


def test_repartition_table_inserted():
    before = len(Document(PATH).tables)
    out = _fill()
    assert len(out.tables) == before + 1, "un tableau (répartition) doit être ajouté"
    rep = next((t for t in out.tables
                if t.rows and "fonctionnalité" in t.rows[0].cells[0].text.strip().lower()), None)
    assert rep is not None, "tableau Fonctionnalité→Composant introuvable"
    # 1 en-tête + 3 lignes de données
    assert len(rep.rows) == 4
    assert rep.rows[1].cells[1].text.strip() == "Sécurité / Backend"


def test_planning_jh_and_total():
    out = _fill()
    activities = _extract_planning_activities(Document(PATH))
    n = len(activities)
    assert n > 0
    table = _find_planning_table(out)
    # TOTAL = somme des J/H écrits (chaque activité reçoit "7")
    total_cell = next((r.cells[-1].text.strip() for r in table.rows
                       if "total" in r.cells[0].text.strip().lower()), None)
    assert total_cell == str(7 * n), f"TOTAL attendu {7 * n}, obtenu {total_cell}"


def test_team_untouched_without_cv():
    base_team = [r.cells[0].text.strip() for r in _find_team_table(Document(PATH)).rows[1:]]
    out = _fill(cvs=[])
    after = [r.cells[0].text.strip() for r in _find_team_table(out).rows[1:]]
    assert after == base_team, "sans CV sélectionné, le tableau équipe doit rester inchangé"


def test_team_dynamic_with_cv():
    out = _fill(cvs=["CV_Jean_Konan.pdf", "CV_Awa_Traore.docx"])
    rows = [(r.cells[0].text.strip(), r.cells[-1].text.strip())
            for r in _find_team_table(out).rows[1:]]
    assert rows == [("Jean Konan", "Chef de Projet"),
                    ("Awa Traore", "Consultant / Développeur Senior")]


def test_cover_date_current_month():
    mois, annee = _cover_date_parts()
    out = _fill()
    boxtext = ""
    for tx in out.element.body.findall(".//" + qn("w:txbxContent")):
        boxtext += "".join((t.text or "") for t in tx.findall(".//" + qn("w:t")))
    assert mois in boxtext and annee in boxtext


if __name__ == "__main__":
    import sys
    sys.exit(pytest.main([__file__, "-v"]))
