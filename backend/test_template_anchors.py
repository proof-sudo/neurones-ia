"""Diagnostic : le template contient-il les ancres attendues par OfferGenerator ?"""
import unicodedata
from docx import Document

PATH = "../data/ged/offres-techniques/TEMPLATE/Offre_Technique_MCI_CARE.docx"

LEGACY = ["BGFI BANK", "BGFI", "KORAZ PARTNERS", "KORAZ", "SOCOPRIM"]
HEADINGS = [
    "EXPRESSION DES BESOINS",
    "OBJECTIFS DE NEURONES TECHNOLOGIES",
    "PRESENTATION DE LA REPONSE DE NEURONES TECHNOLOGIES",
    "Description détaillée de la solution proposée",
    "MÉTHODOLOGIE DE TRAVAIL ET DESCRIPTION DES SERVICES",
]


def norm(s):
    s = unicodedata.normalize("NFC", s)
    return (s.replace("’", "'").replace(" ", " ").replace("–", "-")
             .replace("—", "-").lower())


doc = Document(PATH)

print("=" * 70)
print("STYLES DE TITRE PRÉSENTS (les paragraphes 'Heading*') :")
print("=" * 70)
heading_styles = set()
all_headings = []
for p in doc.paragraphs:
    name = p.style.name if p.style else "?"
    if name.startswith("Heading") and p.text.strip():
        heading_styles.add(name)
        all_headings.append((name, p.text.strip()))
if heading_styles:
    print("Styles trouvés :", sorted(heading_styles))
    print(f"\n{len(all_headings)} headings 'Heading*' :")
    for st, txt in all_headings[:40]:
        print(f"  [{st}] {txt[:70]}")
else:
    print("⚠️  AUCUN paragraphe avec un style commençant par 'Heading' !")
    print("    → _find_heading_idx ne trouvera JAMAIS de section → tout est skip.")
    # montrer les styles réellement utilisés
    used = {}
    for p in doc.paragraphs:
        if p.text.strip():
            used[p.style.name] = used.get(p.style.name, 0) + 1
    print("\nStyles réellement utilisés dans le doc :")
    for st, n in sorted(used.items(), key=lambda x: -x[1]):
        print(f"  {n:4d}× {st}")

print("\n" + "=" * 70)
print("ANCRES DE SECTIONS (recherchées dans TOUT le texte, pas que headings) :")
print("=" * 70)
full_norm = norm("\n".join(p.text for p in doc.paragraphs))
for h in HEADINGS:
    found = norm(h) in full_norm
    print(f"  {'OK ' if found else 'ABSENT'} : {h}")

print("\n" + "=" * 70)
print("NOMS LEGACY (BGFI / KORAZ / SOCOPRIM) :")
print("=" * 70)
raw = "\n".join(p.text for p in doc.paragraphs)
for lg in LEGACY:
    print(f"  {'OK ' if lg in raw else 'ABSENT'} : {lg}")

print("\n" + "=" * 70)
print(f"TABLES : {len(doc.tables)} (le code attend >= 2 ; planning = dernière)")
print("=" * 70)

print("\n" + "=" * 70)
print("STYLE RÉEL DES PARAGRAPHES PORTANT CHAQUE ANCRE DE SECTION :")
print("=" * 70)
for h in HEADINGS:
    frag = norm(h)
    hits = [(i, p.style.name if p.style else "?", p.text.strip())
            for i, p in enumerate(doc.paragraphs)
            if frag in norm(p.text)]
    if hits:
        for i, st, txt in hits:
            print(f"  [{st}]  (para #{i})  «{txt[:60]}»")
    else:
        print(f"  ABSENT des paragraphes du corps : {h}  (peut-être dans une table ?)")
