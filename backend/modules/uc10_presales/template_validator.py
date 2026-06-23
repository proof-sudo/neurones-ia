"""
Validateur (préflight) de template d'offre `.docx` — UC10 Pre-Sales.

Vérifie qu'un `.docx` respecte le contrat attendu par `offer_generator` AVANT de
lancer une génération réelle, et renvoie un rapport ✅/❌ par exigence indiquant
précisément quoi corriger dans le `.docx`.

Principe clé : on RÉUTILISE les helpers de détection du générateur
(`_norm`, `_find_heading_idx`, `_find_template`) — donc ce que le rapport annonce
est exactement ce que le générateur fera réellement (pas de logique dupliquée qui
pourrait diverger).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument

from modules.uc10_presales.template_contract import (
    DEFAULT_DOMAIN, REQUIRED_HEADINGS, PLANNING_PHASES, LEGACY_NAMES,
    MIN_TABLES, TECH_TABLE_DATA_ROWS, PLANNING_MIN_PHASES_MATCH,
    PH_TITRE, PH_CLIENT, PH_MOIS, PH_ANNEE, PH_BESOINS, PH_BESOINS_TYPO,
    PH_PROPOSITION, PH_FONCTIONNALITES, PH_REPARTITION,
)
from modules.uc10_presales.offer_generator import (
    _norm, _find_heading_idx, _find_template, _find_planning_table,
    _ph_norm, _placeholders_present, _present_placeholder_keys,
    _find_team_table, _extract_planning_activities,
)

# Sévérités. Un « error » invalide le template (génération cassée / section vide).
# Un « warning » dégrade le résultat mais laisse la génération produire un document.
SEV_ERROR = "error"
SEV_WARNING = "warning"

# Nombre de paragraphes du début scannés pour la page de garde (cf. _cover_paragraphs).
_COVER_SCAN = 40


@dataclass
class TemplateCheck:
    label: str
    ok: bool
    detail: str = ""
    severity: str = SEV_ERROR


@dataclass
class TemplateValidation:
    ok: bool
    domain: str
    template_path: str | None
    errors: int
    warnings: int
    checks: list[TemplateCheck] = field(default_factory=list)


def make_validation(domain: str, template_path: str | None,
                    checks: list[TemplateCheck]) -> TemplateValidation:
    """Construit le rapport agrégé : `ok` est vrai s'il ne reste aucune erreur."""
    errors = sum(1 for c in checks if c.severity == SEV_ERROR and not c.ok)
    warnings = sum(1 for c in checks if c.severity == SEV_WARNING and not c.ok)
    return TemplateValidation(
        ok=errors == 0, domain=domain, template_path=template_path,
        errors=errors, warnings=warnings, checks=checks,
    )


# Placeholders attendus + sévérité. Un placeholder « contenu » absent = section non
# remplie → erreur ; date = avertissement ; répartition = avertissement (tableau non inséré).
_EXPECTED_PH = [
    (PH_TITRE, [PH_TITRE], SEV_ERROR),
    (PH_CLIENT, [PH_CLIENT], SEV_ERROR),
    (PH_MOIS, [PH_MOIS], SEV_WARNING),
    (PH_ANNEE, [PH_ANNEE], SEV_WARNING),
    (PH_BESOINS, [PH_BESOINS, PH_BESOINS_TYPO], SEV_ERROR),
    (PH_PROPOSITION, [PH_PROPOSITION], SEV_ERROR),
    (PH_FONCTIONNALITES, [PH_FONCTIONNALITES], SEV_ERROR),
    (PH_REPARTITION, [PH_REPARTITION], SEV_WARNING),
]


def _validate_placeholder_document(doc) -> list[TemplateCheck]:
    """Valide un template « à variables » {{ }} : présence des placeholders attendus,
    tableau équipe (NOM/FONCTION) et tableau de planning avec activités détectables."""
    checks: list[TemplateCheck] = []
    present = _present_placeholder_keys(doc)

    for label, variants, sev in _EXPECTED_PH:
        found = any(_ph_norm(v) in present for v in variants)
        checks.append(TemplateCheck(
            f"Placeholder « {{{{{label}}}}} »", found,
            "présent" if found else
            "introuvable → la valeur correspondante ne sera pas insérée", sev))

    team = _find_team_table(doc)
    checks.append(TemplateCheck(
        "Tableau équipe (NOM / FONCTION)", team is not None,
        "trouvé" if team is not None else
        "introuvable (en-tête NOM + FONCTION/RÔLE) → équipe non alimentée par les CV",
        SEV_WARNING))

    acts = _extract_planning_activities(doc)
    checks.append(TemplateCheck(
        "Tableau de planning (activités)", bool(acts),
        f"{len(acts)} activité(s) détectée(s)" if acts else
        "aucune activité détectée → J/H non mis à jour", SEV_WARNING))

    return checks


# ── Cœur : valider un Document déjà ouvert ────────────────────────────────────

def validate_open_document(doc) -> list[TemplateCheck]:
    """Aiguille vers la validation adaptée au type de template :
    - « à variables » {{ }} → contrat placeholders ;
    - sinon → contrat legacy (noms legacy + ancres de titres + 3 tableaux)."""
    if _placeholders_present(doc):
        return _validate_placeholder_document(doc)
    return _validate_legacy_document(doc)


def _validate_legacy_document(doc) -> list[TemplateCheck]:
    paras = doc.paragraphs
    checks: list[TemplateCheck] = []

    # 1. Ancres de titres : présence ET ordre.
    last_idx = -1
    for h in REQUIRED_HEADINGS:
        idx = _find_heading_idx(paras, h)
        if idx is None:
            checks.append(TemplateCheck(
                f"Heading « {h} »", False,
                "introuvable (style Heading*) → section non remplie", SEV_ERROR))
        elif idx <= last_idx:
            checks.append(TemplateCheck(
                f"Heading « {h} »", False,
                f"présent (paragraphe #{idx}) mais hors ordre — attendu après #{last_idx}",
                SEV_ERROR))
            last_idx = max(last_idx, idx)
        else:
            checks.append(TemplateCheck(
                f"Heading « {h} »", True, f"trouvé (paragraphe #{idx})", SEV_ERROR))
            last_idx = idx

    # 2. Nom legacy en page de garde : sans lui, titre/client/date non remplacés.
    cover = _norm(" ".join(p.text for p in paras[:_COVER_SCAN]))
    legacy_hit = next((n for n in LEGACY_NAMES if _norm(n) in cover), None)
    checks.append(TemplateCheck(
        "Nom legacy en page de garde", legacy_hit is not None,
        f"« {legacy_hit} » détecté" if legacy_hit else
        "aucun nom legacy (KORAZ / BGFI / SOCOPRIM…) dans les premiers paragraphes "
        "→ titre / client / date ne seront pas remplacés",
        SEV_WARNING))

    # 3. Nombre de tableaux : composants (0), équipe (1), planning (-1).
    n_tables = len(doc.tables)
    checks.append(TemplateCheck(
        f"Au moins {MIN_TABLES} tableaux", n_tables >= MIN_TABLES,
        f"{n_tables} tableau(x) trouvé(s) — requis : composants, équipe, planning",
        SEV_ERROR))

    # 4. Tableau composants (1er) : en-tête + lignes de stack.
    if n_tables >= 1:
        t0 = doc.tables[0]
        ok0 = len(t0.columns) >= 2 and len(t0.rows) >= 2
        checks.append(TemplateCheck(
            "Tableau composants (1er)", ok0,
            f"{len(t0.rows)} ligne(s) × {len(t0.columns)} colonne(s) — requis : "
            f"1 en-tête + jusqu'à {TECH_TABLE_DATA_ROWS} lignes de stack, ≥ 2 colonnes",
            SEV_WARNING))

    # 5. Tableau équipe (2e) : ≥ 2 colonnes (Rôle / Mission).
    if n_tables >= 2:
        t1 = doc.tables[1]
        checks.append(TemplateCheck(
            "Tableau équipe (2e)", len(t1.columns) >= 2,
            f"{len(t1.columns)} colonne(s) — requis : ≥ 2 (Rôle, Mission)",
            SEV_WARNING))

    # 6. Tableau planning : localisé par CONTENU (même logique que le générateur),
    #    puis vérification de la présence de chaque phase.
    planning = _find_planning_table(doc) if n_tables >= 1 else None
    if planning is None:
        checks.append(TemplateCheck(
            "Tableau de planning", False,
            f"aucun tableau ne contient ≥ {PLANNING_MIN_PHASES_MATCH} phases connues "
            "→ le planning (J/H) ne sera pas mis à jour",
            SEV_WARNING))
    else:
        ptxt = _norm(" ".join(c.text for r in planning.rows for c in r.cells))
        for ph in PLANNING_PHASES:
            found = _norm(ph) in ptxt
            checks.append(TemplateCheck(
                f"Phase planning « {ph} »", found,
                "présente dans le tableau de planning" if found else
                "introuvable dans le tableau de planning → J/H de cette phase non mis à jour",
                SEV_WARNING))

    return checks


# ── Wrappers ──────────────────────────────────────────────────────────────────

def validate_template(source) -> list[TemplateCheck]:
    """Valide un `.docx`. `source` : chemin (str/Path) ou objet fichier (BytesIO)."""
    doc = DocxDocument(str(source) if isinstance(source, (str, Path)) else source)
    return validate_open_document(doc)


def validate_domain_template(domain: str = DEFAULT_DOMAIN) -> TemplateValidation:
    """Valide le template présent dans la GED pour `domain` (via `_find_template`)."""
    path = _find_template(domain)
    if path is None:
        check = TemplateCheck(
            "Template .docx présent", False,
            f"aucun .docx dans data/ged/offres-techniques/TEMPLATE/{domain}/ "
            "(ni dans le dossier TEMPLATE de repli)",
            SEV_ERROR)
        return make_validation(domain, None, [check])
    return make_validation(domain, str(path), validate_template(path))


# ── Rendu texte (CLI) ─────────────────────────────────────────────────────────

def format_report(v: TemplateValidation) -> str:
    lines = [
        f"Template : {v.template_path or '(introuvable)'}",
        f"Domaine  : {v.domain}",
        f"Résultat : {'✅ VALIDE' if v.ok else '❌ INVALIDE'}  "
        f"({v.errors} erreur(s), {v.warnings} avertissement(s))",
        "",
    ]
    for c in v.checks:
        icon = "✅" if c.ok else ("❌" if c.severity == SEV_ERROR else "⚠️")
        suffix = f" — {c.detail}" if (c.detail and not c.ok) else ""
        lines.append(f"{icon} {c.label}{suffix}")
    return "\n".join(lines)


if __name__ == "__main__":
    # Usage : python -m modules.uc10_presales.template_validator [<chemin.docx>|<domaine>]
    import sys

    # Console Windows souvent en cp1252 → l'emoji du rapport (✅/❌/⚠️) ferait planter print().
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    arg = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_DOMAIN
    candidate = Path(arg)
    if candidate.suffix.lower() == ".docx" and candidate.exists():
        report = make_validation("(fichier)", str(candidate), validate_template(candidate))
    else:
        report = validate_domain_template(arg)
    print(format_report(report))
    sys.exit(0 if report.ok else 1)
