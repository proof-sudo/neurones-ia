"""
Export de la matrice de conformité en Excel (UC10 Pre-Sales).

VUE générée depuis `MatriceConformite` (source de vérité) — jamais une source de
données. Deux feuilles :
  - « Matrice de conformité » : une ligne par exigence, en-tête figé + auto-filtre,
    listes déroulantes (Domaine validé, Statut) pour la validation humaine ;
  - « Synthèse » : répartition par domaine / statut / type + contrôle d'exhaustivité.

Anti-perte : on écrit TOUTES les exigences de la matrice, dans l'ordre.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.domain.requirements import (
    MatriceConformite,
    DOMAINE_NON_CLASSE,
    STATUT_A_TRAITER, STATUT_CONFORME, STATUT_CONFORME_PARTIEL,
    STATUT_NON_CONFORME, STATUT_NON_APPLICABLE,
)
from modules.uc10_presales.template_contract import DOMAINS

# ── Libellés humains (le code interne reste le vocabulaire contrôlé) ──────────
_STATUT_LABELS = {
    STATUT_A_TRAITER: "À traiter",
    STATUT_CONFORME: "Conforme",
    STATUT_CONFORME_PARTIEL: "Partiellement conforme",
    STATUT_NON_CONFORME: "Non conforme",
    STATUT_NON_APPLICABLE: "Non applicable",
}
_TYPE_LABELS = {
    "BESOIN": "Besoin", "CRITERE": "Critère", "PREREQUIS": "Prérequis",
    "RESSOURCE": "Ressource", "VIGILANCE": "Vigilance", "PROFIL": "Profil",
    "SEUIL": "Seuil", "ANNEXE": "Annexe", "GRILLE": "Grille",
}

_NAVY = "0F295A"
_BLUE = "1E40AF"
_GRAY_BG = "F1F5F9"
_RED = "B91C1C"

_HEADERS = [
    ("Réf.", 10),
    ("Type", 12),
    ("Exigence", 70),
    ("Source AO", 16),
    ("Domaine(s) suggéré(s)", 24),
    ("Domaine validé", 18),
    ("Section réponse", 20),
    ("Statut", 22),
    ("Bloquant", 10),
    ("Commentaire", 30),
]

_THIN = Side(style="thin", color="D0D7E2")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _statut_label(code: str) -> str:
    return _STATUT_LABELS.get(code, _STATUT_LABELS[STATUT_A_TRAITER])


def render_matrix_xlsx(matrice: MatriceConformite) -> bytes:
    """Sérialise la matrice de conformité en classeur .xlsx (bytes)."""
    wb = Workbook()
    _build_matrix_sheet(wb.active, matrice)
    _build_summary_sheet(wb.create_sheet("Synthèse"), matrice)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_matrix_sheet(ws, matrice: MatriceConformite) -> None:
    ws.title = "Matrice de conformité"

    header_fill = PatternFill("solid", fgColor=_NAVY)
    header_font = Font(bold=True, color="FFFFFF", size=10)
    for col, (label, width) in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 28

    for i, ex in enumerate(matrice.exigences):
        r = i + 2
        values = [
            ex.id,
            _TYPE_LABELS.get(ex.type, ex.type),
            ex.texte,
            ex.source_ref or "—",
            ", ".join(ex.domaines_suggeres) or DOMAINE_NON_CLASSE,
            ex.domaine_valide,
            ex.section_reponse,
            _statut_label(ex.statut_conformite),
            "OUI" if ex.blocking else "",
            ex.commentaire,
        ]
        zebra = "FFFFFF" if i % 2 == 0 else "F8FAFC"
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.border = _BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=col in (3, 5, 7, 10),         # exigence, domaines, section, commentaire
                horizontal="center" if col in (1, 2, 9) else "left",
            )
            cell.fill = PatternFill("solid", fgColor=zebra)
            if col == 3:
                cell.font = Font(size=10)
            if col == 9 and ex.blocking:                # bloquant en rouge gras
                cell.font = Font(bold=True, color=_RED)
                cell.fill = PatternFill("solid", fgColor="FEE2E2")

    n = len(matrice.exigences)
    last_row = n + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_HEADERS))}{max(last_row, 1)}"

    # Listes déroulantes pour la validation humaine (sur la plage de données).
    if n:
        rng = f"{last_row}"
        dv_dom = DataValidation(
            type="list",
            formula1='"' + ",".join([""] + list(DOMAINS) + [DOMAINE_NON_CLASSE]) + '"',
            allow_blank=True,
        )
        dv_statut = DataValidation(
            type="list",
            formula1='"' + ",".join(_STATUT_LABELS.values()) + '"',
            allow_blank=True,
        )
        ws.add_data_validation(dv_dom)
        ws.add_data_validation(dv_statut)
        dv_dom.add(f"F2:F{rng}")        # Domaine validé
        dv_statut.add(f"H2:H{rng}")     # Statut


def _build_summary_sheet(ws, matrice: MatriceConformite) -> None:
    title_font = Font(bold=True, color=_NAVY, size=13)
    head_font = Font(bold=True, color="FFFFFF", size=10)
    head_fill = PatternFill("solid", fgColor=_BLUE)

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 12

    ws["A1"] = "Synthèse — Matrice de conformité"
    ws["A1"].font = title_font

    meta = [
        ("Appel d'offres", matrice.ao_filename),
        ("Généré le", matrice.generated_at or "—"),
        ("Total exigences", matrice.total),
        ("Attendu (contrôle)", matrice.expected_total if matrice.expected_total is not None else "—"),
        ("Exhaustif", "OUI" if matrice.is_exhaustive else "NON — exigences perdues !"),
    ]
    row = 3
    for label, val in meta:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=row, column=2, value=val)
        if label == "Exhaustif" and not matrice.is_exhaustive:
            c.font = Font(bold=True, color=_RED)
        row += 1

    def _table(title: str, counts: dict, label_map: dict | None = None) -> None:
        nonlocal row
        row += 1
        ws.cell(row=row, column=1, value=title).font = title_font
        row += 1
        for col, lbl in enumerate(("Catégorie", "Nombre"), start=1):
            c = ws.cell(row=row, column=col, value=lbl)
            c.font = head_font
            c.fill = head_fill
        row += 1
        for key, val in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])):
            ws.cell(row=row, column=1, value=(label_map or {}).get(key, key))
            ws.cell(row=row, column=2, value=val)
            row += 1

    _table("Par domaine", matrice.count_by_domaine())
    _table("Par statut", matrice.count_by_statut(), _STATUT_LABELS)
    _table("Par type", matrice.count_by_type(), _TYPE_LABELS)
