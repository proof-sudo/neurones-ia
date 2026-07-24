"""
Construction d'une offre technique Word **de zéro**, sans aucun template .docx.

L'IA ne « bouche » plus des trous de gabarit : elle produit un PLAN DE DOCUMENT
(une couverture + une liste de blocs typés), et ce module met ce plan en forme dans
un .docx à la charte Neurones (orange #FF6600), en construisant tout par le code :

  • styles (police, couleurs, titres H1/H2/H3, style de tableau) ;
  • page de garde : logo + titre/sous-titre/accroche (rédigés par l'IA) + client + date ;
  • en-tête (logo réduit) et pied de page (numéro de page) ;
  • rendu de chaque bloc IA (heading / paragraph / bullets / numbered / table / page_break) ;
  • blocs INSTITUTIONNELS cannés côté code (présentation Neurones, méthodologie,
    certifications) — des faits sur l'entreprise, jamais laissés à l'invention du LLM ;
  • tableau de l'équipe projet alimenté par les CV choisis dans le modal ;
  • annexes : CV / ABE choisis, chaque page PDF rendue en image pleine page.

Vocabulaire de blocs accepté (tout autre `type` est ignoré silencieusement) :
  {"type": "heading",   "level": 1|2|3, "text": "..."}
  {"type": "paragraph", "text": "..."}
  {"type": "bullets",   "items": ["...", ...]}
  {"type": "numbered",  "items": ["...", ...]}
  {"type": "table",     "titre": "...", "headers": ["..."], "rows": [["..."], ...]}
  {"type": "page_break"}
"""

import logging
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor, Inches

from config.settings import settings

logger = logging.getLogger(__name__)

# ── Charte Neurones Technologies ──────────────────────────────────────────────
PRIMARY = RGBColor(0xFF, 0x66, 0x00)        # orange vif — filets, fond d'en-tête de tableau
PRIMARY_DARK = RGBColor(0xC2, 0x4E, 0x00)   # orange foncé — titres (contraste AA sur blanc)
INK = RGBColor(0x1F, 0x29, 0x37)            # gris ardoise — corps de texte
MUTED = RGBColor(0x6B, 0x72, 0x80)          # gris — accroche / mentions discrètes
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_PRIMARY_HEX = "FF6600"
_BODY_FONT = "Calibri"

# Logo backend (copié depuis frontend/public). Résolu par rapport à la racine du
# paquet backend → robuste quel que soit le cwd (dev local, conteneur).
_LOGO_PATH = Path(__file__).resolve().parents[2] / "assets" / "logo.png"

_MONTHS_FR = {
    1: "Janvier", 2: "Février", 3: "Mars", 4: "Avril",
    5: "Mai", 6: "Juin", 7: "Juillet", 8: "Août",
    9: "Septembre", 10: "Octobre", 11: "Novembre", 12: "Décembre",
}


def _cover_date_parts() -> tuple[str, str]:
    """(« Mois », « Année ») du jour, pour la page de garde et le nom de fichier."""
    now = datetime.now()
    return _MONTHS_FR[now.month], str(now.year)


# ── Styles du document ────────────────────────────────────────────────────────

def _configure_styles(doc: DocxDocument) -> None:
    """Applique la charte aux styles de base du document (corps + titres)."""
    normal = doc.styles["Normal"]
    normal.font.name = _BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color in (
        ("Heading 1", 17, PRIMARY_DARK),
        ("Heading 2", 13.5, PRIMARY_DARK),
        ("Heading 3", 11.5, INK),
    ):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = _BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True


def _set_cell_shading(cell, hex_fill: str) -> None:
    """Remplit le fond d'une cellule (w:shd dans w:tcPr)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _add_heading_rule(doc: DocxDocument, para) -> None:
    """Ajoute un filet orange sous un titre de niveau 1 (repère visuel de section)."""
    p_pr = para._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), _PRIMARY_HEX)
    borders.append(bottom)
    p_pr.append(borders)


# ── En-tête / pied de page ────────────────────────────────────────────────────

def _add_page_number_field(paragraph) -> None:
    """Insère un champ « PAGE » (numéro de page dynamique) dans le paragraphe."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _configure_header_footer(doc: DocxDocument, client_name: str) -> None:
    """En-tête : logo réduit + libellé ; pied : mention + numéro de page.
    Non affichés sur la page de garde (first-page header/footer distincts et vides)."""
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if _LOGO_PATH.exists():
        hp.add_run().add_picture(str(_LOGO_PATH), height=Inches(0.28))
    else:
        r = hp.add_run("NEURONES TECHNOLOGIES")
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = PRIMARY_DARK

    footer = section.footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = f"Offre technique — {client_name}  ·  " if client_name else "Offre technique  ·  "
    r = fp.add_run(label)
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    _add_page_number_field(fp)


# ── Page de garde ─────────────────────────────────────────────────────────────

def _build_cover(doc: DocxDocument, cover: dict, client_name: str) -> None:
    """Page de garde : logo + titre / sous-titre / accroche (rédigés par l'IA),
    puis client et date. Terminée par un saut de page."""
    mois, annee = _cover_date_parts()

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(60)

    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if _LOGO_PATH.exists():
        logo_p.add_run().add_picture(str(_LOGO_PATH), width=Inches(2.4))
    else:
        r = logo_p.add_run("NEURONES TECHNOLOGIES")
        r.bold = True
        r.font.size = Pt(20)
        r.font.color.rgb = PRIMARY_DARK

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = Pt(50)
    kr = kicker.add_run((cover.get("sous_titre") or "OFFRE TECHNIQUE").upper())
    kr.bold = True
    kr.font.size = Pt(12)
    kr.font.color.rgb = PRIMARY
    _spread_letters(kr)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    tr = title.add_run(cover.get("titre") or "Offre technique")
    tr.bold = True
    tr.font.size = Pt(26)
    tr.font.color.rgb = PRIMARY_DARK

    accroche = (cover.get("accroche") or "").strip()
    if accroche:
        ap = doc.add_paragraph()
        ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ap.paragraph_format.space_before = Pt(14)
        ar = ap.add_run(accroche)
        ar.italic = True
        ar.font.size = Pt(12)
        ar.font.color.rgb = MUTED

    if client_name:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(60)
        cr = cp.add_run("Préparée pour")
        cr.font.size = Pt(10)
        cr.font.color.rgb = MUTED
        cn = doc.add_paragraph()
        cn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cnr = cn.add_run(client_name)
        cnr.bold = True
        cnr.font.size = Pt(15)
        cnr.font.color.rgb = INK

    dp = doc.add_paragraph()
    dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dp.paragraph_format.space_before = Pt(24)
    dr = dp.add_run(f"{mois} {annee}")
    dr.font.size = Pt(11)
    dr.font.color.rgb = MUTED

    doc.add_page_break()


def _spread_letters(run) -> None:
    """Léger espacement des lettres (effet « kicker ») via w:spacing sur le run."""
    rpr = run._r.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), "40")  # twips
    rpr.append(spacing)


# ── Rendu des blocs IA ────────────────────────────────────────────────────────

def _clean_text(value: object) -> str:
    """Texte affichable d'une valeur de bloc (tolère None / nombres)."""
    return str(value if value is not None else "").strip()


def _render_heading(doc: DocxDocument, block: dict) -> None:
    try:
        level = int(block.get("level", 1))
    except (TypeError, ValueError):
        level = 1
    level = min(max(level, 1), 3)
    text = _clean_text(block.get("text"))
    if not text:
        return
    p = doc.add_heading(text, level=level)
    if level == 1:
        _add_heading_rule(doc, p)


def _render_paragraph(doc: DocxDocument, block: dict) -> None:
    text = _clean_text(block.get("text"))
    if text:
        doc.add_paragraph(text)


def _render_list(doc: DocxDocument, block: dict, numbered: bool) -> None:
    style = "List Number" if numbered else "List Bullet"
    resolved = style if style in doc.styles else "Normal"
    for item in block.get("items", []) or []:
        text = _clean_text(item)
        if not text:
            continue
        p = doc.add_paragraph(text, style=resolved)
        if resolved == "Normal":  # pas de style de liste dans le doc → puce manuelle
            p.text = f"•  {text}"


def _render_table(doc: DocxDocument, block: dict) -> None:
    headers = [_clean_text(h) for h in (block.get("headers") or [])]
    raw_rows = block.get("rows") or []
    rows = [[_clean_text(c) for c in (r or [])] for r in raw_rows if isinstance(r, list)]
    n_cols = max([len(headers)] + [len(r) for r in rows] or [0])
    if n_cols == 0:
        return

    titre = _clean_text(block.get("titre"))
    if titre:
        tp = doc.add_paragraph()
        tr = tp.add_run(titre)
        tr.bold = True
        tr.font.color.rgb = INK

    table = doc.add_table(rows=0, cols=n_cols)
    if "Table Grid" in [s.name for s in doc.styles]:
        table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    if headers:
        cells = table.add_row().cells
        for i in range(n_cols):
            _set_cell_shading(cells[i], _PRIMARY_HEX)
            cell = cells[i]
            cell.text = headers[i] if i < len(headers) else ""
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = WHITE
    for r in rows:
        cells = table.add_row().cells
        for i in range(n_cols):
            cells[i].text = r[i] if i < len(r) else ""

    doc.add_paragraph()  # respiration après le tableau


_BLOCK_RENDERERS = {
    "heading": _render_heading,
    "paragraph": _render_paragraph,
    "bullets": lambda doc, b: _render_list(doc, b, numbered=False),
    "numbered": lambda doc, b: _render_list(doc, b, numbered=True),
    "table": _render_table,
    "page_break": lambda doc, b: doc.add_page_break(),
}


def _render_blocks(doc: DocxDocument, blocks: list[dict]) -> None:
    """Rend chaque bloc via son renderer ; les types inconnus sont ignorés."""
    for block in blocks or []:
        if not isinstance(block, dict):
            continue
        renderer = _BLOCK_RENDERERS.get(str(block.get("type", "")).strip().lower())
        if renderer is not None:
            renderer(doc, block)


# ── Blocs institutionnels cannés (faits Neurones — jamais générés par le LLM) ──
# Prose neutre et professionnelle, SANS chiffres/certifications/références inventés.
# À personnaliser une fois avec les éléments réels de l'entreprise.

def presentation_neurones_blocks() -> list[dict]:
    return [
        {"type": "heading", "level": 1, "text": "Présentation de Neurones Technologies"},
        {"type": "paragraph", "text":
            "Neurones Technologies est une société de services et d'ingénierie informatique "
            "qui accompagne les organisations dans leur transformation numérique. Nos équipes "
            "conçoivent, développent et déploient des solutions logicielles, des infrastructures "
            "et des dispositifs de cybersécurité adaptés aux enjeux métiers de nos clients."},
        {"type": "paragraph", "text":
            "Notre démarche associe expertise technique, rigueur méthodologique et proximité "
            "avec les équipes du client, afin de garantir des livrables de qualité, dans les "
            "délais et le budget convenus."},
        {"type": "bullets", "items": [
            "Développement d'applications sur mesure (web, mobile, back-office).",
            "Infrastructure, réseaux et hébergement (on-premise et cloud).",
            "Cybersécurité : audit, sécurisation et supervision.",
            "Conseil, intégration et accompagnement au changement.",
        ]},
    ]


def methodologie_blocks() -> list[dict]:
    return [
        {"type": "heading", "level": 1, "text": "Méthodologie et gestion de projet"},
        {"type": "paragraph", "text":
            "Neurones Technologies structure ses projets selon une approche par phases, "
            "jalonnée de livrables et de points de validation avec le client. Cette démarche "
            "assure la maîtrise des délais, de la qualité et des risques tout au long du projet."},
        {"type": "bullets", "items": [
            "Cadrage : ateliers de recueil des besoins, validation du périmètre et du planning.",
            "Conception : architecture technique et fonctionnelle, spécifications détaillées.",
            "Réalisation : développement itératif et paramétrage de la solution.",
            "Recette : plan de tests, validation fonctionnelle et corrections.",
            "Déploiement : mise en production et documentation d'exploitation.",
            "Formation et transfert de compétences aux équipes du client.",
            "Garantie et support post-déploiement.",
        ]},
        {"type": "paragraph", "text":
            "Un chef de projet dédié assure la coordination, le reporting régulier et le suivi "
            "des risques, en interlocuteur unique du client pendant toute la durée du projet."},
    ]


def certifications_blocks() -> list[dict]:
    return [
        {"type": "heading", "level": 1, "text": "Certifications et engagements qualité"},
        {"type": "paragraph", "text":
            "Nos consultants entretiennent leurs compétences sur les principales technologies "
            "du marché (éditeurs logiciels, plateformes cloud, solutions de cybersécurité et "
            "référentiels de gestion de projet), gage de la fiabilité de nos réalisations."},
        {"type": "bullets", "items": [
            "Respect des standards et bonnes pratiques de développement et de sécurité.",
            "Documentation complète et transfert de compétences systématiques.",
            "Engagement sur les délais, la qualité et la confidentialité des données.",
            "Accompagnement et support après la mise en production.",
        ]},
    ]


# ── Tableau de l'équipe projet (alimenté par les CV choisis) ──────────────────

def _clean_person_name(filename: str) -> str:
    """« CV_Jean_Konan.docx » → « Jean Konan » : retire extension, préfixe CV, séparateurs."""
    name = filename.rsplit(".", 1)[0]
    name = re.sub(r"^cv[\s_-]+", "", name, flags=re.IGNORECASE)
    return name.replace("_", " ").replace("-", " ").strip()


def _named_team_from_cvs(
    selected_cvs: list[str], kb_team: dict[str, tuple[str, str]] | None = None
) -> list[tuple[str, str]]:
    """(Nom, Fonction) à partir des CV choisis.

    Priorité à `kb_team` (nom_complet/titre_poste RÉELS, extraits du CV et
    persistés dans `kb_cv` par l'indexation GED) — le nom de fichier n'est
    qu'un repli pour les CV non encore résolus dans la base de connaissance
    (jamais de "Chef de Projet" arbitraire sur le 1er de la liste si une
    vraie fonction est disponible)."""
    kb_team = kb_team or {}
    people: list[tuple[str, str]] = []
    for i, filename in enumerate(f for f in selected_cvs if f):
        resolved = kb_team.get(filename)
        if resolved and resolved[0]:
            nom = resolved[0]
            fonction = resolved[1] or ("Chef de Projet" if i == 0 else "Consultant / Développeur Senior")
        else:
            nom = _clean_person_name(filename).strip().title()
            fonction = "Chef de Projet" if i == 0 else "Consultant / Développeur Senior"
        people.append((nom, fonction))
    return people


def _build_team_section(
    doc: DocxDocument, selected_cvs: list[str], kb_team: dict[str, tuple[str, str]] | None = None
) -> None:
    """Section « Équipe projet dédiée » : tableau NOM | FONCTION depuis les CV choisis.
    Sans CV, un tableau générique de rôles est produit (repli)."""
    people = _named_team_from_cvs(selected_cvs, kb_team)
    if not people:
        people = [
            ("Chef de Projet", "Coordination et pilotage du projet"),
            ("Expert Technique Senior", "Architecture et développement"),
            ("Développeur Senior", "Implémentation et intégration"),
            ("Ingénieur QA / Recette", "Tests, validation et documentation"),
        ]
    _render_heading(doc, {"level": 1, "text": "Équipe projet dédiée"})
    doc.add_paragraph(
        "Neurones Technologies mobilise pour ce projet une équipe pluridisciplinaire "
        "aux compétences complémentaires :"
    )
    _render_table(doc, {
        "headers": ["Nom", "Fonction"],
        "rows": [[name, func] for name, func in people],
    })


# ── Annexes CV / ABE (pages PDF → images pleine page) ─────────────────────────

_TODO_COLOR = RGBColor(0xC0, 0x00, 0x00)


def _resolve_ged_files(filenames: list[str], folder: Path) -> list[Path]:
    """Mappe des noms de fichiers vers leur chemin réel sous la GED (récursif).
    Les introuvables sont ignorés silencieusement — la génération n'est jamais bloquée."""
    if not filenames or not folder.exists():
        return []
    by_name = {f.name: f for f in folder.rglob("*") if f.is_file()}
    return [by_name[n] for n in filenames if n in by_name]


def _insert_pdf_pages(doc: DocxDocument, pdf_path: Path, max_pages: int = 25) -> tuple[int, int]:
    """Insère chaque page du PDF comme image pleine page (aspect préservé, centrée).
    Retourne (pages_insérées, pages_totales)."""
    import fitz

    sec = doc.sections[0]
    usable_w = int(sec.page_width - sec.left_margin - sec.right_margin)
    usable_h = int(sec.page_height - sec.top_margin - sec.bottom_margin)
    mat = fitz.Matrix(130 / 72.0, 130 / 72.0)  # ~130 DPI : lisible sans alourdir

    inserted, total = 0, 0
    with fitz.open(str(pdf_path)) as pdf:
        total = pdf.page_count
        n_to_insert = min(max_pages, total)
        for i, page in enumerate(pdf):
            if i >= max_pages:
                break
            try:
                pix = page.get_pixmap(matrix=mat)
                png = pix.tobytes("png")
                ratio = (pix.width / pix.height) if pix.height else 1.0
                target_w = usable_w
                if usable_h and (target_w / ratio) > usable_h:
                    target_w = int(usable_h * ratio)
                doc.add_picture(BytesIO(png), width=Emu(target_w))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                inserted += 1
                if i < n_to_insert - 1:
                    doc.add_page_break()
            except Exception as exc:
                logger.warning("Page %d de %s non insérée (%s)", i + 1, pdf_path.name, exc)
    return inserted, total


def _append_annexes(doc: DocxDocument, cv_paths: list[Path], abe_paths: list[Path]) -> None:
    """Intègre en fin d'offre les CV et ABE choisis : chaque page PDF en image pleine page.
    Les formats non-PDF sont signalés « à annexer manuellement ». Sans sélection, ne fait rien."""
    groups: list[tuple[str, list[Path]]] = []
    if cv_paths:
        groups.append(("ANNEXE — CV DES INTERVENANTS", cv_paths))
    if abe_paths:
        groups.append(("ANNEXE — ATTESTATIONS DE BONNE EXÉCUTION (ABE)", abe_paths))
    if not groups:
        return

    for heading, paths in groups:
        doc.add_page_break()
        p = doc.add_heading(heading, level=1)
        _add_heading_rule(doc, p)
        for path in paths:
            sub = doc.add_paragraph()
            sr = sub.add_run(path.stem.replace("_", " ").replace("-", " ").strip())
            sr.bold = True
            sr.font.color.rgb = PRIMARY_DARK
            if path.suffix.lower() == ".pdf":
                inserted, total = _insert_pdf_pages(doc, path)
                if inserted == 0:
                    nr = doc.add_paragraph().add_run(
                        f"⟦À COMPLÉTER⟧ Impossible d'extraire « {path.name} » — à annexer manuellement.")
                    nr.font.color.rgb = _TODO_COLOR
                elif total > inserted:
                    nr = doc.add_paragraph().add_run(
                        f"⟦À COMPLÉTER⟧ {total - inserted} page(s) de « {path.name} » non insérée(s) "
                        f"(limite {inserted}).")
                    nr.font.color.rgb = _TODO_COLOR
            else:
                nr = doc.add_paragraph().add_run(
                    f"⟦À COMPLÉTER⟧ « {path.name} » (format {path.suffix}) à annexer manuellement.")
                nr.font.color.rgb = _TODO_COLOR
            doc.add_page_break()


# ── Point d'entrée ────────────────────────────────────────────────────────────

def build_offer_docx(
    cover: dict,
    blocks: list[dict],
    client_name: str,
    selected_cvs: list[str] | None = None,
    selected_abes: list[str] | None = None,
    kb_team: dict[str, tuple[str, str]] | None = None,
) -> bytes:
    """Construit le .docx complet à partir du plan de document produit par l'IA.

    Assemblage (les blocs cannés et l'équipe sont ajoutés par le code, l'IA ne les
    produit pas) :
        Couverture (IA) → Présentation Neurones → Corps technique (IA)
        → Méthodologie → Équipe projet → Certifications → Annexes CV/ABE
    """
    client_name = (client_name or "").strip()
    doc = DocxDocument()
    _configure_styles(doc)
    _configure_header_footer(doc, client_name)

    _build_cover(doc, cover or {}, client_name)
    _render_blocks(doc, presentation_neurones_blocks())
    _render_blocks(doc, blocks or [])
    _render_blocks(doc, methodologie_blocks())
    _build_team_section(doc, selected_cvs or [], kb_team)
    _render_blocks(doc, certifications_blocks())

    ged_root = Path(settings.ged_path)
    cv_paths = _resolve_ged_files(selected_cvs or [], ged_root / "cvs")
    abe_paths = _resolve_ged_files(selected_abes or [], ged_root / "abe")
    logger.info(
        "Annexes offre — CV résolus=%d/%d, ABE résolus=%d/%d (ged=%s)",
        len(cv_paths), len(selected_cvs or []), len(abe_paths), len(selected_abes or []), ged_root,
    )
    _append_annexes(doc, cv_paths, abe_paths)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
