import asyncio
import hashlib
import json
import logging
import re
import unicodedata
import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from config.settings import settings

from fastapi import APIRouter, Request, UploadFile, File, Form, Body, HTTPException
from fastapi.responses import Response, StreamingResponse

from modules.uc10_presales.schemas import (
    ScoringResultSchema, OfferGenerationRequest, KeyElementSchema, MatchedDocumentSchema, BidRecommendationSchema,
    BidStrategyRequest, BidStrategySchema, AnalysisExportRequest,
    StrategyExportRequest, ChecklistExportRequest, MatrixConfirmRequest,
    TeamMatchRequest, TeamMatchResponse,
    MarketIdentitySchema, CalendarEventSchema, EvaluationModalitiesSchema,
    ScoringCriterionSchema, RiskSchema, PreconditionSchema, AppendixSchema,
    PartnerSchema, PhaseActionSchema, StrategyPhaseSchema,
    RequiredProfileSchema, EligibilityThresholdSchema, FinancialDataSchema,
    CapabilityDealSchema, CapabilityMatchSchema, ClientContextSchema,
    OfferSectionsSchema, OfferSectionsResponse, OfferRenderRequest,
)
from modules.uc10_presales.use_case import PresalesUseCase
from modules.uc10_presales.offer_generator import OfferGenerator
from core.domain.offer import ScoringResult, BidStrategy, Partner, Appendix
# Matrice de conformité : build_matrice / render_matrix_xlsx étaient UTILISÉS sans être
# importés (NameError latent sur /export-matrix) → import explicite. + assesseur IA,
# store de persistance (clôture C5) et utilitaires de latence.
from modules.uc10_presales.requirements_builder import build_matrice
from modules.uc10_presales.matrix_export import render_matrix_xlsx
from modules.uc10_presales.conformity import assess_matrice
from modules.uc10_presales import matrix_store
from modules.uc10_presales import dossier_store
from modules.uc10_presales.latency import prune_cache_dir, content_key

# Checklist générique inférée (CI / marchés publics) — utilisée quand l'AO ne liste
# aucune pièce explicite. Affichée avec un avertissement « inférée ».
_GENERIC_CHECKLIST: list[Appendix] = [
    Appendix(code="—", label="Registre du commerce (RCCM)", type="ADMIN"),
    Appendix(code="—", label="Attestation de régularité fiscale (DGI)", type="ADMIN"),
    Appendix(code="—", label="Attestation de régularité sociale (CNPS)", type="ADMIN"),
    Appendix(code="—", label="Statuts de la société et pouvoir du signataire", type="ADMIN"),
    Appendix(code="—", label="Attestation de non-faillite / non-exclusion", type="ADMIN"),
    Appendix(code="—", label="Références similaires avec attestations de bonne fin", type="TECHNIQUE"),
    Appendix(code="—", label="Méthodologie d'intervention et moyens techniques", type="TECHNIQUE"),
    Appendix(code="—", label="Bilans des 3 derniers exercices", type="FINANCIER"),
    Appendix(code="—", label="Caution bancaire de soumission", type="FINANCIER"),
    Appendix(code="—", label="CVs des profils clés et certifications", type="RH"),
]

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/presales", tags=["UC10 - Pre-Sales"])

_ALLOWED_TYPES = {"application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _attachment_headers(filename: str) -> dict[str, str]:
    """Content-Disposition sûr : Starlette encode les en-têtes HTTP en latin-1, donc un
    caractère hors latin-1 dans le nom (—, ', …) fait crasher Response() en 500. Repli
    ASCII pour filename= + nom UTF-8 complet en filename* (RFC 5987) pour le navigateur."""
    ascii_name = (
        unicodedata.normalize("NFKD", filename)
        .encode("ascii", "ignore")
        .decode()
        .replace('"', "'")
        .strip()
    ) or "document.docx"
    return {
        "Content-Disposition": f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"
    }


def _get_use_case(request: Request) -> PresalesUseCase:
    container = request.app.state.container
    return PresalesUseCase(
        llm_haiku=container.llm_haiku,
        llm_sonnet=container.llm_sonnet,
        rag_engine=container.rag_engine,
        pdf_parser=container.pdf_parser,
        docx_parser=container.docx_parser,
        vision_ocr=container.presales_vision_ocr,
    )


# Cache d'analyse sur DISQUE (pas de dépendance externe). Fichier JSON nommé par le SHA-256
# du contenu de l'AO → ré-analyser le MÊME document renvoie le MÊME résultat (le pipeline LLM
# n'est pas déterministe, même à température 0 : c'est la seule garantie de reproductibilité).
# `force=true` recalcule et écrase l'entrée. Content-addressé → pas de TTL nécessaire.
_SCORE_CACHE_DIR = Path(settings.uploads_path).parent / "score_cache"
# Cache optionnel des sections d'offre (clé = empreinte scoring + client + modèle).
_OFFER_CACHE_DIR = Path(settings.uploads_path).parent / "offer_sections_cache"


def _score_cache_path(file_bytes: bytes) -> Path:
    return _SCORE_CACHE_DIR / f"{hashlib.sha256(file_bytes).hexdigest()}.json"


# Intervalle du « heartbeat » pendant le scoring long (~169s en prod). Le pipeline n'émet
# rien avant sa réponse finale ; sur une connexion inactive, un proxy/pare-feu côté client
# coupe à ~30s (→ 499 nginx, 504 navigateur). On streame donc un espace toutes les 10s : des
# espaces en tête d'un JSON sont valides, donc `response.json()` côté client les ignore et
# aucun changement de parsing n'est requis pour le cas nominal.
_SCORE_HEARTBEAT_SECONDS = 10


@router.post("/score")
async def score_ao(
    request: Request,
    file: UploadFile = File(..., description="AO en PDF ou Word"),
    force: bool = False,
):
    """Upload un AO et lance le pipeline de scoring en 5 étapes.

    Le cache disque est désactivé par défaut en prod (settings.score_cache_enabled) pour ne pas
    saturer le serveur : chaque appel recalcule. S'il est réactivé, le résultat est mémorisé par
    empreinte SHA-256 du fichier et `?force=true` force une nouvelle analyse.

    Réponse streamée (`application/json`) : des espaces de heartbeat maintiennent la connexion
    active pendant le pipeline (cf. `_SCORE_HEARTBEAT_SECONDS`), suivis du JSON du résultat.
    Comme le flux démarre par un `200`, une erreur du pipeline ne peut plus être un code HTTP :
    elle est encodée dans le corps avec une sentinelle `{"__error__": <status>, "detail": ...}`
    que le client détecte.
    """
    logger.info("Score AO reçu — fichier=%s content_type=%s force=%s", file.filename, file.content_type, force)
    if file.content_type not in _ALLOWED_TYPES and not file.filename.endswith((".pdf", ".docx")):
        raise HTTPException(status_code=400, detail="Format non supporté. Utilisez PDF ou DOCX.")

    file_bytes = await file.read()
    if len(file_bytes) > settings.presales_max_upload_mb * 1024 * 1024:
        raise HTTPException(
            status_code=400,
            detail=f"Fichier trop volumineux (max {settings.presales_max_upload_mb} Mo).",
        )

    # Cache disque désactivé en prod (settings.score_cache_enabled=False) : évite de saturer
    # le disque du serveur. Quand actif, sert le résultat mémorisé par empreinte SHA-256.
    cache_path = _score_cache_path(file_bytes) if settings.score_cache_enabled else None
    if settings.score_cache_enabled and cache_path.exists() and not force:
        try:
            cached = json.loads(cache_path.read_text(encoding="utf-8"))
            schema = ScoringResultSchema.model_validate(cached)
            logger.info("Score AO servi depuis le cache disque — %s", file.filename)
            # Cache : réponse immédiate, pas besoin de streamer.
            return Response(
                content=json.dumps(schema.model_dump(mode="json"), ensure_ascii=False),
                media_type="application/json",
            )
        except Exception as exc:
            logger.warning("Cache score illisible (%s) — recalcul de %s", exc, file.filename)

    use_case = _get_use_case(request)

    async def _stream():
        task = asyncio.create_task(use_case.score_ao(filename=file.filename, file_bytes=file_bytes))
        # Heartbeat : espaces réguliers tant que le pipeline tourne (shield → wait_for
        # n'annule pas la tâche en cas de timeout, il continue en arrière-plan).
        while not task.done():
            try:
                await asyncio.wait_for(asyncio.shield(task), timeout=_SCORE_HEARTBEAT_SECONDS)
            except asyncio.TimeoutError:
                yield b" "

        try:
            result = task.result()
        except ValueError as e:
            yield json.dumps({"__error__": 422, "detail": str(e)}, ensure_ascii=False).encode()
            return
        except Exception as e:
            logger.exception("Erreur pipeline scoring AO '%s'", file.filename)
            yield json.dumps(
                {"__error__": 500, "detail": f"Erreur analyse AO : {type(e).__name__}: {e}"},
                ensure_ascii=False,
            ).encode()
            return

        schema = _to_schema(result)
        if settings.score_cache_enabled:
            try:
                _SCORE_CACHE_DIR.mkdir(parents=True, exist_ok=True)
                cache_path.write_text(json.dumps(schema.model_dump(mode="json"), ensure_ascii=False), encoding="utf-8")
                prune_cache_dir(_SCORE_CACHE_DIR, settings.presales_cache_max_files)  # purge LRU bornée
            except OSError as exc:
                logger.debug("Écriture cache score échouée (%s) — non bloquant", exc)
        yield json.dumps(schema.model_dump(mode="json"), ensure_ascii=False).encode()

    return StreamingResponse(_stream(), media_type="application/json")


# ── Dossiers persistés (fichier + état complet du workflow) ──────────────────────
#
# Avant : le fichier AO et tout l'état (décision, stratégie, checklist...) ne
# vivaient qu'en mémoire navigateur (localStorage + File en RAM) → après un
# rechargement de page, « refaire une étape » échouait faute de fichier. Ici, le
# fichier est écrit sur disque et l'état complet (forme `AOEntry` du frontend,
# blob JSON opaque pour le backend) est répliqué en base à chaque changement.
# Purge automatique : cf. `dossier_store.purge_expired` + job planifié.


@router.post("/dossiers")
async def create_dossier(
    file: UploadFile = File(...),
    state: str = Form(..., description="État initial AOEntry (JSON) — id/filename/clientName/owner/deadline/status inclus"),
    dossier_id: str | None = Form(None),
):
    """Crée un dossier persisté : écrit le fichier AO sur disque + sauvegarde l'état initial.

    Le frontend reste propriétaire du schéma `AOEntry` (id généré côté client, état initial
    construit par `newEntry()`) — le backend ne fait que stocker et rejouer ce blob JSON.
    """
    if file.content_type not in _ALLOWED_TYPES and not file.filename.endswith((".pdf", ".docx")):
        raise HTTPException(status_code=400, detail="Format non supporté. Utilisez PDF ou DOCX.")
    file_bytes = await file.read()
    if len(file_bytes) > settings.presales_max_upload_mb * 1024 * 1024:
        raise HTTPException(
            status_code=400,
            detail=f"Fichier trop volumineux (max {settings.presales_max_upload_mb} Mo).",
        )
    try:
        state_dict = json.loads(state)
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="État initial invalide (JSON attendu).")

    dossier_id = dossier_id or state_dict.get("id") or str(uuid.uuid4())
    state_dict["id"] = dossier_id
    file_path = dossier_store.save_file(dossier_id, file.filename, file_bytes)
    saved = await dossier_store.create(dossier_id, file.filename, file_path, state_dict)
    logger.info("Dossier présale créé — id=%s fichier=%s", dossier_id, file.filename)
    return saved


@router.get("/dossiers")
async def list_dossiers():
    """Historique complet des dossiers présale (page d'accueil du module)."""
    return await dossier_store.list_all()


@router.get("/dossiers/{dossier_id}")
async def get_dossier(dossier_id: str):
    state = await dossier_store.get(dossier_id)
    if state is None:
        raise HTTPException(status_code=404, detail="Dossier introuvable.")
    return state


@router.get("/dossiers/{dossier_id}/file")
async def download_dossier_file(dossier_id: str):
    """Télécharge le fichier AO original tel qu'uploadé (le bouton téléchargement de l'historique)."""
    meta = await dossier_store.get_file_meta(dossier_id)
    if meta is None:
        raise HTTPException(status_code=404, detail="Dossier introuvable.")
    file_path, filename = meta
    file_bytes = dossier_store.read_file(file_path)
    if file_bytes is None:
        raise HTTPException(status_code=404, detail="Fichier source introuvable sur le serveur.")
    media_type = _DOCX_MIME if filename.lower().endswith(".docx") else "application/pdf"
    return Response(content=file_bytes, media_type=media_type, headers=_attachment_headers(filename))


@router.patch("/dossiers/{dossier_id}")
async def patch_dossier(dossier_id: str, changes: dict = Body(...)):
    """Fusionne un patch partiel (mêmes clés que l'objet `AOEntry` frontend) dans l'état persisté."""
    updated = await dossier_store.patch(dossier_id, changes)
    if updated is None:
        raise HTTPException(status_code=404, detail="Dossier introuvable.")
    return updated


@router.delete("/dossiers/{dossier_id}")
async def delete_dossier(dossier_id: str):
    ok = await dossier_store.delete(dossier_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Dossier introuvable.")
    return {"deleted": True}


# Analyses en cours : référence FORTE aux tâches détachées. asyncio ne conserve qu'une
# weakref sur les tâches créées par create_task ; sans ce set, une analyse longue peut
# être ramassée par le GC en plein vol. On retire la tâche via un done_callback.
_ANALYSIS_TASKS: set[asyncio.Task] = set()


async def _run_analysis_task(
    dossier_id: str, filename: str, file_bytes: bytes, use_case: PresalesUseCase,
) -> None:
    """Exécute le pipeline de scoring EN ARRIÈRE-PLAN et persiste l'issue dans le dossier
    (`status` → `scored` + `scoringResult`, ou `error` + `errorMessage`). Découplé du cycle
    de vie de la requête HTTP : le client déclenche via `POST .../analyze` (202 immédiat)
    puis interroge `GET /dossiers/{id}` jusqu'à `scored`/`error`.

    Pourquoi : le pipeline dure ~2 min. En streaming, Next.js bufferisait la réponse (le
    heartbeat n'atteignait jamais le proxy) et le frontal Traefik coupait à ~100s → 504.
    Sans connexion longue, plus de 504 possible."""
    try:
        result = await use_case.score_ao(filename=filename, file_bytes=file_bytes)
    except ValueError as e:
        await dossier_store.patch(dossier_id, {"status": "error", "errorMessage": str(e)})
        return
    except Exception as e:
        logger.exception("Erreur pipeline scoring AO (dossier %s)", dossier_id)
        await dossier_store.patch(dossier_id, {
            "status": "error",
            "errorMessage": f"Erreur analyse AO : {type(e).__name__}: {e}",
        })
        return

    schema = _to_schema(result)
    await dossier_store.patch(dossier_id, {
        "status": "scored",
        "scoringResult": schema.model_dump(mode="json"),
        "errorMessage": None,
    })


@router.post("/dossiers/{dossier_id}/analyze", status_code=202)
async def analyze_dossier(dossier_id: str, request: Request, force: bool = False):
    """(Re)lance le pipeline de scoring sur le fichier PERSISTÉ du dossier — permet de
    « refaire l'analyse » à tout moment, y compris après un rechargement de page (plus
    besoin que le File soit encore en mémoire navigateur).

    Modèle ASYNCHRONE : le scoring part en tâche de fond et l'endpoint répond **202
    immédiatement** avec `{"status": "scoring", "dossier_id": ...}`. Le client interroge
    ensuite `GET /dossiers/{id}` jusqu'à ce que `status` passe à `scored` (résultat dans
    `scoringResult`) ou `error` (message dans `errorMessage`). Remplace l'ancien streaming
    heartbeat qui provoquait des 504 (cf. `_run_analysis_task`)."""
    file_path = await dossier_store.get_file_path(dossier_id)
    if file_path is None:
        raise HTTPException(status_code=404, detail="Dossier introuvable.")
    file_bytes = dossier_store.read_file(file_path)
    if file_bytes is None:
        await dossier_store.patch(dossier_id, {
            "status": "error",
            "errorMessage": "Fichier source introuvable sur le serveur — supprimez ce dossier et re-déposez le fichier.",
        })
        raise HTTPException(status_code=422, detail="Fichier source introuvable sur le serveur.")

    current = await dossier_store.get(dossier_id)
    filename = (current or {}).get("filename") or Path(file_path).name

    await dossier_store.patch(dossier_id, {"status": "scoring", "errorMessage": None})
    use_case = _get_use_case(request)

    task = asyncio.create_task(_run_analysis_task(dossier_id, filename, file_bytes, use_case))
    _ANALYSIS_TASKS.add(task)
    task.add_done_callback(_ANALYSIS_TASKS.discard)

    return {"status": "scoring", "dossier_id": dossier_id}


@router.post("/generate")
async def generate_offer(body: OfferGenerationRequest, request: Request):
    """Génère une offre technique Word à partir du résultat de scoring."""
    use_case = _get_use_case(request)
    scoring = _from_schema(body.scoring_result)
    draft = await use_case.generate_offer(
        scoring=scoring,
        client_name=body.client_name or "",
        selected_cvs=body.selected_cvs,
        selected_abes=body.selected_abes,
    )
    return Response(
        content=draft.content_docx,
        media_type=_DOCX_MIME,
        headers=_attachment_headers(draft.filename),
    )


@router.post("/offer/sections", response_model=OfferSectionsResponse)
async def offer_sections(body: OfferGenerationRequest, request: Request):
    """Étape 1 : génère (IA) les sections éditables de l'offre, sans produire le .docx."""
    # Cache disque optionnel (presales_strategy_cache_enabled) : même AO + client + modèle →
    # mêmes sections sans nouvel appel LLM (la génération est non déterministe à T=0.7).
    cache_path = None
    if settings.presales_strategy_cache_enabled:
        key = content_key(body.scoring_result.model_dump_json(), body.client_name or "",
                          settings.offer_sections_model)
        cache_path = _OFFER_CACHE_DIR / f"{key}.json"
        if cache_path.exists():
            try:
                return OfferSectionsResponse(**json.loads(cache_path.read_text(encoding="utf-8")))
            except Exception as exc:
                logger.warning("Cache sections illisible (%s) — régénération", exc)

    use_case = _get_use_case(request)
    scoring = _from_schema(body.scoring_result)
    try:
        sections, domain, client = await use_case.build_offer_sections(scoring, body.client_name or "")
    except Exception as e:
        logger.exception("Génération des sections d'offre échouée")
        raise HTTPException(status_code=500, detail=f"{type(e).__name__}: {e}")
    resp = OfferSectionsResponse(
        sections=OfferSectionsSchema(**sections),
        domain=domain,
        client_name=client,
        filename=OfferGenerator.build_filename(scoring, client),
    )
    if cache_path is not None:
        try:
            _OFFER_CACHE_DIR.mkdir(parents=True, exist_ok=True)
            cache_path.write_text(resp.model_dump_json(), encoding="utf-8")
            prune_cache_dir(_OFFER_CACHE_DIR, settings.presales_cache_max_files)
        except OSError as exc:
            logger.debug("Écriture cache sections échouée (%s) — non bloquant", exc)
    return resp


@router.post("/offer/render")
async def offer_render(body: OfferRenderRequest, request: Request):
    """Étape 2 : produit le .docx Word à partir des sections (éventuellement éditées)."""
    use_case = _get_use_case(request)
    scoring = _from_schema(body.scoring_result)
    logger.info(
        "Offer render — CV sélectionnés=%s | ABE sélectionnés=%s",
        body.selected_cvs, body.selected_abes,
    )
    try:
        draft = await use_case.render_offer(
            scoring, body.sections.model_dump(), body.client_name or "",
            selected_cvs=body.selected_cvs, selected_abes=body.selected_abes,
        )
    except FileNotFoundError as e:
        raise HTTPException(status_code=422, detail=str(e))
    except Exception as e:
        logger.exception("Rendu de l'offre échoué")
        raise HTTPException(status_code=500, detail=f"{type(e).__name__}: {e}")
    return Response(
        content=draft.content_docx,
        media_type=_DOCX_MIME,
        headers=_attachment_headers(draft.filename),
    )


@router.post("/bid-strategy", response_model=BidStrategySchema)
async def bid_strategy(body: BidStrategyRequest, request: Request):
    """Génère une stratégie de réponse structurée (5 phases + étape 0 + appendices)."""
    use_case = _get_use_case(request)
    scoring = _from_schema(body.scoring_result)
    partner = Partner(**body.partner.model_dump()) if body.partner else None
    try:
        result = await use_case.generate_bid_strategy(
            scoring=scoring,
            client_name=body.client_name or "",
            decision=body.decision,
            decision_reason=body.decision_reason or "",
            partner=partner,
        )
        return _strategy_to_schema(result)
    except Exception as e:
        logger.exception("Bid strategy generation failed")
        raise HTTPException(status_code=500, detail=str(e))


def _strategy_to_schema(s: BidStrategy) -> BidStrategySchema:
    return BidStrategySchema(
        phases=[
            StrategyPhaseSchema(
                id=p.id, name=p.name, description=p.description,
                start_day=p.start_day, end_day=p.end_day,
                actions=[PhaseActionSchema(**a.__dict__) for a in p.actions],
                prerequisites=p.prerequisites, is_blocking_next=p.is_blocking_next,
            )
            for p in s.phases
        ],
        strategy_text=s.strategy_text,
        response_plan=s.response_plan,
        appendices=[AppendixSchema(**a.__dict__) for a in s.appendices],
        partner=PartnerSchema(**s.partner.__dict__) if s.partner else None,
        partner_validation=[PreconditionSchema(**pv.__dict__) for pv in s.partner_validation],
        version=s.version,
        parent_version=s.parent_version,
        generated_at=s.generated_at,
    )


def _docx_helpers(doc):
    """Retourne les helpers de mise en forme partagés entre les deux exports."""
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    C_NAVY  = RGBColor(0x0F, 0x29, 0x5A)
    C_BLUE  = RGBColor(0x1E, 0x40, 0xAF)
    C_GREEN = RGBColor(0x15, 0x80, 0x3D)
    C_RED   = RGBColor(0xB9, 0x1C, 0x1C)
    C_AMBER = RGBColor(0x92, 0x40, 0x0E)
    C_GRAY  = RGBColor(0x64, 0x74, 0x8B)
    C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    C_TEXT  = RGBColor(0x1E, 0x29, 0x3B)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    def _cell_bg(cell, hex_color: str):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def _cell_style(cell, font_size=10, bold=False, color=None, bg=None):
        if bg:
            _cell_bg(cell, bg)
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.size = Pt(font_size)
                if bold:
                    run.bold = True
                if color:
                    run.font.color.rgb = color

    def _section_title(text: str, color: RGBColor = None, num: str = ""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run((f"{num}. " if num else "") + text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = color or C_BLUE
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "1E40AF")
        pBdr.append(bottom)
        pPr.append(pBdr)
        return p

    def _bullet(text: str, color: RGBColor = None):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        b = p.add_run("▸  ")
        b.font.color.rgb = color or C_BLUE
        b.font.size = Pt(9)
        r = p.add_run(text.strip())
        r.font.size = Pt(10)
        r.font.color.rgb = C_TEXT
        return p

    def _body_para(text: str, justify=True):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        if justify:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text.strip())
        r.font.size = Pt(10)
        r.font.color.rgb = C_TEXT
        return p

    def _render_rich_text(text: str):
        """Convertit du texte avec markdown basique en paragraphes Word bien formatés."""
        # Normalise les patterns inline "(1) … (2) … (3) …" en liste numérotée
        # multi-lignes pour que chaque item s'affiche sur sa propre ligne.
        if re.search(r"\(\d+\)", text):
            text = re.sub(r"\s*\((\d+)\)\s*", lambda m: f"\n\n{m.group(1)}. ", text)
        paragraphs = re.split(r"\n{2,}", text.strip())
        for block in paragraphs:
            lines = [l.strip() for l in block.split("\n") if l.strip()]
            for line in lines:
                # Puce markdown ou tiret
                if re.match(r"^[-•*]\s+", line):
                    content = re.sub(r"^[-•*]\s+", "", line)
                    _bullet(_strip_bold(content))
                # Liste numérotée
                elif re.match(r"^\d+[.)]\s+", line):
                    content = re.sub(r"^\d+[.)]\s+", "", line)
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.6)
                    p.paragraph_format.space_after = Pt(3)
                    m = re.match(r"^(\d+)", line)
                    num_r = p.add_run(f"{m.group(1)}.  ")
                    num_r.font.color.rgb = C_BLUE
                    num_r.bold = True
                    num_r.font.size = Pt(9)
                    txt_r = p.add_run(_strip_bold(content))
                    txt_r.font.size = Pt(10)
                    txt_r.font.color.rgb = C_TEXT
                # Ligne courte en majuscules = sous-titre implicite
                elif line.isupper() and len(line) < 80:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.space_after = Pt(3)
                    r = p.add_run(line)
                    r.bold = True
                    r.font.size = Pt(10)
                    r.font.color.rgb = C_NAVY
                else:
                    # Paragraphe avec gras inline **...**
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(5)
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    parts = re.split(r"\*\*(.+?)\*\*", line)
                    for idx, part in enumerate(parts):
                        if not part:
                            continue
                        r = p.add_run(part)
                        r.font.size = Pt(10)
                        r.font.color.rgb = C_TEXT
                        if idx % 2 == 1:
                            r.bold = True

    def _strip_bold(text: str) -> str:
        return re.sub(r"\*\*(.+?)\*\*", r"\1", text)

    def _cover_table(rows_data: list[tuple], col_widths=(Cm(5.5), Cm(10.5))):
        t = doc.add_table(rows=len(rows_data), cols=2)
        t.style = "Table Grid"
        for i, (label, value) in enumerate(rows_data):
            row = t.rows[i]
            row.cells[0].text = label
            row.cells[1].text = value
            _cell_bg(row.cells[0], "0F295A")
            _cell_bg(row.cells[1], "EFF6FF" if i % 2 == 0 else "F8FAFC")
            for para in row.cells[0].paragraphs:
                for run in para.runs:
                    run.font.color.rgb = C_WHITE
                    run.font.size = Pt(10)
                    run.bold = True
            for para in row.cells[1].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = C_TEXT
            row.cells[0].width = col_widths[0]
            row.cells[1].width = col_widths[1]
        return t

    def _footer_line():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"Neurones Technologies CI  ·  Document confidentiel  ·  Généré le {datetime.now().strftime('%d/%m/%Y')}")
        r.font.size = Pt(8)
        r.font.color.rgb = C_GRAY
        r.italic = True

    return dict(
        C_NAVY=C_NAVY, C_BLUE=C_BLUE, C_GREEN=C_GREEN,
        C_RED=C_RED, C_AMBER=C_AMBER, C_GRAY=C_GRAY,
        C_WHITE=C_WHITE, C_TEXT=C_TEXT,
        cell_bg=_cell_bg, cell_style=_cell_style,
        section_title=_section_title, bullet=_bullet,
        body_para=_body_para, render_rich_text=_render_rich_text,
        cover_table=_cover_table, footer_line=_footer_line,
        Pt=Pt, RGBColor=RGBColor, Cm=Cm,
        WD_ALIGN_PARAGRAPH=WD_ALIGN_PARAGRAPH,
    )


@router.post("/export-analysis")
async def export_analysis(body: AnalysisExportRequest, request: Request):
    """Exporte l'analyse complète de l'AO en document Word professionnel."""
    from docx import Document as DocxDocument

    scoring = _from_schema(body.scoring_result)
    client_name = body.client_name or "Non précisé"
    doc = DocxDocument()
    h = _docx_helpers(doc)
    Pt = h["Pt"]; Cm = h["Cm"]
    WD = h["WD_ALIGN_PARAGRAPH"]

    # ── PAGE DE COUVERTURE ────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.alignment = WD.CENTER
    title_p.paragraph_format.space_before = Pt(30)
    title_p.paragraph_format.space_after = Pt(2)
    tr = title_p.add_run("ANALYSE")
    tr.bold = True; tr.font.size = Pt(32); tr.font.color.rgb = h["C_NAVY"]

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD.CENTER
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(4)
    sr = sub_p.add_run("APPEL D'OFFRES")
    sr.bold = True; sr.font.size = Pt(20); sr.font.color.rgb = h["C_BLUE"]

    sep = doc.add_paragraph()
    sep.alignment = WD.CENTER
    sep.paragraph_format.space_after = Pt(30)
    sep.add_run("─" * 45).font.color.rgb = h["C_BLUE"]

    h["cover_table"]([
        ("Appel d'offres", scoring.ao_filename),
        ("Client / Commanditaire", client_name),
        ("Date d'analyse", datetime.now().strftime("%d %B %Y")),
        ("Statut document", "Analyse AO — Confidentiel"),
    ])

    doc.add_page_break()

    section_num = 0
    def _next_num() -> str:
        nonlocal section_num
        section_num += 1
        return str(section_num)

    # ── FICHE D'IDENTITÉ DU MARCHÉ ────────────────────────────────────────────
    identity = scoring.market_identity
    identity_rows = [
        ("Type de marché", identity.type_marche),
        ("Référence", identity.reference),
        ("Autorité contractante", identity.autorite_contractante),
        ("Durée du contrat", identity.duree_contrat),
        ("Date de démarrage", identity.date_demarrage),
        ("Deadline de soumission", identity.deadline_soumission),
        ("Validité de l'offre", identity.validite_offre),
        ("Périmètre géographique", identity.perimetre_geographique),
        ("Éligibilité candidat", identity.eligibilite_candidat),
    ]
    identity_rows = [(lbl, val) for lbl, val in identity_rows if val and val.strip()]
    if identity_rows:
        h["section_title"]("Fiche d'identité du marché", num=_next_num())
        conf_p = doc.add_paragraph()
        conf_p.paragraph_format.space_after = Pt(4)
        conf_r = conf_p.add_run(f"Fiabilité de l'extraction : {int(identity.confidence * 100)}%")
        conf_r.italic = True; conf_r.font.size = Pt(9); conf_r.font.color.rgb = h["C_GRAY"]
        h["cover_table"](identity_rows)
        doc.add_paragraph("")

    # ── RÉSUMÉ EXÉCUTIF ───────────────────────────────────────────────────────
    h["section_title"]("Résumé exécutif", num=_next_num())
    if scoring.summary:
        h["body_para"](scoring.summary)

    # ── CALENDRIER DE L'AO ────────────────────────────────────────────────────
    if scoring.calendar:
        h["section_title"]("Calendrier de l'AO", num=_next_num())
        cal_t = doc.add_table(rows=1, cols=3)
        cal_t.style = "Table Grid"
        for cell, lbl in zip(cal_t.rows[0].cells, ["Criticité", "Événement", "Date"]):
            cell.text = lbl
            h["cell_bg"](cell, "1E40AF")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(10)
        crit_color = {"BLOQUANT": "FEE2E2", "CRITIQUE": "FEF3C7", "INFO": "F1F5F9"}
        for i, ev in enumerate(scoring.calendar):
            row = cal_t.add_row().cells
            row[0].text = ev.criticite
            label_text = ev.label + (f"\n({ev.source_section})" if ev.source_section else "")
            row[1].text = label_text
            row[2].text = ev.date
            h["cell_bg"](row[0], crit_color.get(ev.criticite, "F1F5F9"))
            h["cell_bg"](row[1], "FFFFFF" if i % 2 == 0 else "F8FAFC")
            h["cell_bg"](row[2], "FFFFFF" if i % 2 == 0 else "F8FAFC")
            for cell in row:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(10)
            if row[0].paragraphs[0].runs:
                row[0].paragraphs[0].runs[0].bold = True
            row[0].width = Cm(3); row[1].width = Cm(9); row[2].width = Cm(4)
        doc.add_paragraph("")

    # ── POINTS CLÉS ───────────────────────────────────────────────────────────
    if scoring.key_elements:
        h["section_title"]("Points clés identifiés par l'IA", num=_next_num())
        kp_t = doc.add_table(rows=1, cols=2)
        kp_t.style = "Table Grid"
        for cell, lbl in zip(kp_t.rows[0].cells, ["Critère", "Valeur"]):
            cell.text = lbl
            h["cell_bg"](cell, "1E40AF")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(10)
        for i, el in enumerate(scoring.key_elements):
            row = kp_t.add_row().cells
            row[0].text = el.category
            row[1].text = el.value
            h["cell_bg"](row[0], "EFF6FF" if i % 2 == 0 else "FFFFFF")
            h["cell_bg"](row[1], "F8FAFC" if i % 2 == 0 else "FFFFFF")
            for cell in row:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(10)
            if row[0].paragraphs[0].runs:
                row[0].paragraphs[0].runs[0].bold = True
            row[0].width = Cm(6); row[1].width = Cm(10)
        doc.add_paragraph("")

    # ── MODALITÉS D'ÉVALUATION ────────────────────────────────────────────────
    evaluation = scoring.evaluation_modalities
    eval_has_content = (
        evaluation.ponderation_technique > 0 or evaluation.ponderation_financiere > 0
        or evaluation.seuil_minimum_technique > 0
        or bool(evaluation.formule_notation_financiere) or bool(evaluation.modalites)
    )
    if eval_has_content:
        h["section_title"]("Modalités d'évaluation", num=_next_num())
        conf_p = doc.add_paragraph()
        conf_p.paragraph_format.space_after = Pt(4)
        conf_r = conf_p.add_run(f"Fiabilité de l'extraction : {int(evaluation.confidence * 100)}%")
        conf_r.italic = True; conf_r.font.size = Pt(9); conf_r.font.color.rgb = h["C_GRAY"]

        eval_t = doc.add_table(rows=1, cols=3)
        eval_t.style = "Table Grid"
        for cell, lbl in zip(
            eval_t.rows[0].cells,
            ["Pondération technique", "Pondération financière", "Seuil min. technique"],
        ):
            cell.text = lbl
            h["cell_bg"](cell, "1E40AF")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(10)
        vals_row = eval_t.add_row().cells
        vals_row[0].text = (
            f"{evaluation.ponderation_technique} %"
            if evaluation.ponderation_technique > 0 else "Non précisé dans l'AO"
        )
        vals_row[1].text = (
            f"{evaluation.ponderation_financiere} %"
            if evaluation.ponderation_financiere > 0 else "Non précisé dans l'AO"
        )
        vals_row[2].text = (
            f"{evaluation.seuil_minimum_technique} / 100"
            if evaluation.seuil_minimum_technique > 0 else "Non précisé dans l'AO"
        )
        eval_present = [
            evaluation.ponderation_technique > 0,
            evaluation.ponderation_financiere > 0,
            evaluation.seuil_minimum_technique > 0,
        ]
        for cell, bg, present in zip(vals_row, ["EFF6FF", "EEF2FF", "F1F5F9"], eval_present):
            h["cell_bg"](cell, bg)
            for para in cell.paragraphs:
                para.alignment = h["WD_ALIGN_PARAGRAPH"].CENTER
                for run in para.runs:
                    if present:
                        run.bold = True; run.font.size = Pt(14); run.font.color.rgb = h["C_NAVY"]
                    else:
                        run.italic = True; run.font.size = Pt(9); run.font.color.rgb = h["C_GRAY"]
        doc.add_paragraph("")

        if evaluation.formule_notation_financiere:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            lbl_r = p.add_run("Formule de notation financière : ")
            lbl_r.bold = True; lbl_r.font.size = Pt(10); lbl_r.font.color.rgb = h["C_NAVY"]
            val_r = p.add_run(evaluation.formule_notation_financiere)
            val_r.font.size = Pt(10); val_r.font.color.rgb = h["C_TEXT"]; val_r.font.name = "Consolas"

        if evaluation.modalites:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run("Modalités complémentaires")
            r.bold = True; r.font.size = Pt(10); r.font.color.rgb = h["C_NAVY"]
            for m in evaluation.modalites:
                h["bullet"](m, h["C_BLUE"])
        doc.add_paragraph("")

    # ── PROFILS DEMANDÉS (effectifs + compétences détaillées) ─────────────────
    if scoring.profils_demandes:
        total_postes = sum(p.quantite for p in scoring.profils_demandes)
        titre = "Profils demandés" + (f" — {total_postes} postes au total" if total_postes else "")
        h["section_title"](titre, color=h["C_NAVY"], num=_next_num())
        prof_t = doc.add_table(rows=1, cols=5)
        prof_t.style = "Table Grid"
        for cell, lbl in zip(prof_t.rows[0].cells, ["Profil", "Qté", "Niveau / Exp.", "Compétences", "Certifications"]):
            cell.text = lbl
            h["cell_bg"](cell, "1E40AF")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(9)
        for i, p in enumerate(scoring.profils_demandes):
            bg = ["EFF6FF", "F8FAFC", "F8FAFC", "F8FAFC", "F8FAFC"] if i % 2 == 0 else ["FFFFFF"] * 5
            row = prof_t.add_row().cells
            # Col 0 : nom (gras) + domaine (gris) + missions (petites lignes)
            row[0].text = ""
            p0 = row[0].paragraphs[0]
            rn = p0.add_run(p.profil); rn.bold = True; rn.font.size = Pt(9); rn.font.color.rgb = h["C_NAVY"]
            if p.domaine:
                rd = p0.add_run(f"  ·  {p.domaine}"); rd.font.size = Pt(8); rd.font.color.rgb = h["C_GRAY"]
            for m in p.missions:
                mp = row[0].add_paragraph(); mp.paragraph_format.space_after = Pt(0)
                rm = mp.add_run(f"› {m}"); rm.font.size = Pt(8); rm.font.color.rgb = h["C_TEXT"]
            # Cols 1-4
            row[1].text = str(p.quantite) if p.quantite else "—"
            row[2].text = " / ".join(x for x in (p.niveau, p.experience_min) if x) or "—"
            row[3].text = " · ".join(p.competences) if p.competences else "—"
            row[4].text = " · ".join(p.certifications) if p.certifications else "—"
            for ci in range(1, 5):
                for para in row[ci].paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9); run.font.color.rgb = h["C_TEXT"]
            row[1].paragraphs[0].alignment = h["WD_ALIGN_PARAGRAPH"].CENTER
            for ci in range(5):
                h["cell_bg"](row[ci], bg[ci])
            row[0].width = Cm(4.0); row[1].width = Cm(1.2); row[2].width = Cm(2.8); row[3].width = Cm(4.4); row[4].width = Cm(3.2)
        doc.add_paragraph("")

    # ── SEUILS D'ÉLIGIBILITÉ (recevabilité chiffrée) ──────────────────────────
    if scoring.seuils_eligibilite:
        h["section_title"]("Seuils d'éligibilité", color=h["C_RED"], num=_next_num())
        seuil_t = doc.add_table(rows=1, cols=4)
        seuil_t.style = "Table Grid"
        for cell, lbl in zip(seuil_t.rows[0].cells, ["Critère d'éligibilité", "Valeur", "Type", "Éliminatoire"]):
            cell.text = lbl
            h["cell_bg"](cell, "1E40AF")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(9)
        for i, s in enumerate(scoring.seuils_eligibilite):
            row = seuil_t.add_row().cells
            row[0].text = s.libelle
            row[1].text = (s.valeur + (f" {s.unite}" if s.unite else "")).strip() or "—"
            row[2].text = s.type
            row[3].text = "OUI" if s.blocking else "non"
            base_bg = "FFFFFF" if i % 2 else "F8FAFC"
            for ci in range(4):
                h["cell_bg"](row[ci], base_bg)
                for para in row[ci].paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9); run.font.color.rgb = h["C_TEXT"]
            if row[1].paragraphs[0].runs:
                row[1].paragraphs[0].runs[0].bold = True
                row[1].paragraphs[0].runs[0].font.color.rgb = h["C_NAVY"]
            if s.blocking and row[3].paragraphs[0].runs:
                row[3].paragraphs[0].runs[0].bold = True
                row[3].paragraphs[0].runs[0].font.color.rgb = h["C_RED"]
                h["cell_bg"](row[3], "FEE2E2")
            row[3].paragraphs[0].alignment = h["WD_ALIGN_PARAGRAPH"].CENTER
            row[0].width = Cm(7.0); row[1].width = Cm(4.0); row[2].width = Cm(3.0); row[3].width = Cm(2.5)
        doc.add_paragraph("")

    # ── DONNÉES FINANCIÈRES ───────────────────────────────────────────────────
    fin = scoring.donnees_financieres
    fin_rows = [(lbl, val) for lbl, val in [
        ("Budget estimé", fin.budget_estime),
        ("Modalités de paiement", fin.modalites_paiement),
        ("Garantie de soumission", fin.garantie_soumission),
        ("Pénalités", fin.penalites),
    ] if val]
    if fin_rows:
        h["section_title"]("Données financières", color=h["C_GREEN"], num=_next_num())
        h["cover_table"](fin_rows)
        doc.add_paragraph("")

    # ── SECTIONS DOSSIER ──────────────────────────────────────────────────────
    sections = [
        ("Critères de sélection", scoring.criteres_selection, h["C_BLUE"]),
        ("Besoins identifiés", scoring.besoins, h["C_BLUE"]),
        ("Prérequis", scoring.prerequis, h["C_NAVY"]),
        ("Ressources demandées", scoring.ressources_demandees, h["C_NAVY"]),
        ("Points de vigilance", scoring.points_vigilance, h["C_AMBER"]),
    ]
    for title, items, color in sections:
        if items:
            h["section_title"](title, color=color, num=_next_num())
            for item in items:
                # items = list[ExtractedItem] : on affiche le texte + la réf source si connue
                txt = item.texte + (f"  ·  réf. {item.source_section}" if item.source_section else "")
                h["bullet"](txt, color)
            doc.add_paragraph("")

    h["footer_line"]()
    buffer = BytesIO()
    doc.save(buffer)
    safe_name = re.sub(r"\.(pdf|docx)$", "", scoring.ao_filename, flags=re.IGNORECASE).replace(" ", "-")
    return Response(
        content=buffer.getvalue(), media_type=_DOCX_MIME,
        headers=_attachment_headers(f"Analyse-AO_{safe_name}.docx"),
    )


@router.post("/export-matrix")
async def export_matrix(body: AnalysisExportRequest, request: Request):
    """Exporte la MATRICE DE CONFORMITÉ (toutes les exigences de l'AO) en Excel.

    Vue générée depuis le ScoringResult : consolidation exhaustive des exigences
    (besoins, critères, prérequis, profils, seuils, annexes…), classées par domaine,
    avec leur référence source. Aucun appel LLM — pur calcul déterministe."""
    scoring = _from_schema(body.scoring_result)
    matrice = build_matrice(scoring, generated_at=datetime.now().isoformat(timespec="seconds"))
    if settings.matrix_assess_on_export:
        # Pré-statut de conformité IA (déterministe, sans appel LLM) + persistance (clôture C5)
        # → l'Excel exporté et la checklist UI partagent la même matrice assessée.
        assess_matrice(matrice, scoring)
        try:
            matrix_store.save(matrice)
            prune_cache_dir(matrix_store.default_base(), settings.presales_cache_max_files)
        except OSError as exc:
            logger.debug("Persistance matrice échouée (%s) — non bloquant", exc)
    logger.info(
        "Export matrice — %s : %d exigences (exhaustif=%s)",
        scoring.ao_filename, matrice.total, matrice.is_exhaustive,
    )
    try:
        xlsx = render_matrix_xlsx(matrice)
    except Exception as e:
        logger.exception("Rendu de la matrice de conformité échoué")
        raise HTTPException(status_code=500, detail=f"{type(e).__name__}: {e}")
    safe_name = re.sub(r"\.(pdf|docx)$", "", scoring.ao_filename, flags=re.IGNORECASE).replace(" ", "-")
    return Response(
        content=xlsx, media_type=_XLSX_MIME,
        headers=_attachment_headers(f"Matrice-Conformite_{safe_name}.xlsx"),
    )


# ── Checklist de conformité validée par l'IA + contrôle humain par cochage ────────
@router.post("/matrix/assess")
async def matrix_assess(body: AnalysisExportRequest, request: Request):
    """Construit la matrice de conformité, PRÉ-STATUE chaque exigence (IA, déterministe,
    sans appel LLM : dérivé du scoring) et persiste le tout. Renvoie la matrice (dict)
    que la checklist UI affichera, ligne par ligne, pour cochage humain de confirmation."""
    scoring = _from_schema(body.scoring_result)
    matrice = build_matrice(scoring, generated_at=datetime.now().isoformat(timespec="seconds"))
    assess_matrice(matrice, scoring)
    try:
        matrix_store.save(matrice)
        prune_cache_dir(matrix_store.default_base(), settings.presales_cache_max_files)
    except OSError as exc:
        logger.warning("Persistance matrice échouée (%s) — non bloquant", exc)
    logger.info("Matrice assessée — %s : %d exigences", scoring.ao_filename, matrice.total)
    return matrice.to_dict()


@router.get("/matrix")
async def matrix_get(ao_filename: str):
    """Recharge la matrice persistée (incluant les validations humaines déjà saisies)
    d'un AO. 404 si aucune matrice n'a encore été assessée pour cet AO."""
    matrice = matrix_store.load(ao_filename)
    if matrice is None:
        raise HTTPException(status_code=404, detail="Aucune matrice persistée — lancez d'abord /matrix/assess.")
    return matrice.to_dict()


@router.post("/matrix/confirm")
async def matrix_confirm(body: MatrixConfirmRequest):
    """Applique les COCHAGES humains de confirmation (statut validé, domaine, commentaire)
    et re-persiste. Chaque exigence cochée est figée : confirmé par / le. C'est le contrôle
    humain qui atteste que les éléments validés sont effectivement réunis."""
    matrice = matrix_store.confirm(
        body.ao_filename,
        [c.model_dump() for c in body.confirmations],
        confirme_par=body.confirme_par,
        now_iso=datetime.now().isoformat(timespec="seconds"),
    )
    if matrice is None:
        raise HTTPException(status_code=404, detail="Aucune matrice à confirmer — lancez d'abord /matrix/assess.")
    return {
        "ao_filename": matrice.ao_filename,
        "total": matrice.total,
        "confirmes": sum(1 for e in matrice.exigences if e.confirme),
        "par_statut_valide": matrice.count_by_statut(),
        "exigences": matrice.to_dict()["exigences"],
    }


@router.post("/export-scoring")
async def export_scoring(body: AnalysisExportRequest, request: Request):
    """Exporte le scoring & positionnement (Step 2) : score, forces, risques, écarts, docs GED."""
    from docx import Document as DocxDocument

    scoring = _from_schema(body.scoring_result)
    client_name = body.client_name or "Non précisé"
    doc = DocxDocument()
    h = _docx_helpers(doc)
    Pt = h["Pt"]; Cm = h["Cm"]
    WD = h["WD_ALIGN_PARAGRAPH"]

    score_color = h["C_GREEN"] if scoring.score >= 70 else (h["C_AMBER"] if scoring.score >= 40 else h["C_RED"])
    rec_map = {
        "GO": ("GO — Répondre à l'AO", h["C_GREEN"]),
        "CONDITIONAL": ("CONDITIONNEL — Sous réserve", h["C_AMBER"]),
        "NO_BID": ("NO-BID — Ne pas répondre", h["C_RED"]),
    }
    rec_label, rec_color = rec_map.get(scoring.recommendation.value, ("—", h["C_AMBER"]))

    # ── PAGE DE COUVERTURE ────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.alignment = WD.CENTER
    title_p.paragraph_format.space_before = Pt(30)
    title_p.paragraph_format.space_after = Pt(2)
    tr = title_p.add_run("SCORING")
    tr.bold = True; tr.font.size = Pt(32); tr.font.color.rgb = h["C_NAVY"]

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD.CENTER
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(4)
    sr = sub_p.add_run("& POSITIONNEMENT")
    sr.bold = True; sr.font.size = Pt(20); sr.font.color.rgb = h["C_BLUE"]

    sep = doc.add_paragraph()
    sep.alignment = WD.CENTER
    sep.paragraph_format.space_after = Pt(20)
    sep.add_run("─" * 45).font.color.rgb = h["C_BLUE"]

    h["cover_table"]([
        ("Appel d'offres", scoring.ao_filename),
        ("Client / Commanditaire", client_name),
        ("Date d'analyse", datetime.now().strftime("%d %B %Y")),
        ("Statut document", "Scoring & Décision — Confidentiel"),
    ])

    score_p = doc.add_paragraph()
    score_p.alignment = WD.CENTER
    score_p.paragraph_format.space_before = Pt(18)
    score_p.paragraph_format.space_after = Pt(4)
    sc_r = score_p.add_run(f"Score GED : {scoring.score} / 100")
    sc_r.bold = True; sc_r.font.size = Pt(22); sc_r.font.color.rgb = score_color

    rec_p = doc.add_paragraph()
    rec_p.alignment = WD.CENTER
    rec_p.paragraph_format.space_after = Pt(30)
    rr = rec_p.add_run(f"Recommandation IA : {rec_label}")
    rr.bold = True; rr.font.size = Pt(14); rr.font.color.rgb = rec_color

    doc.add_page_break()

    # Numérotation dynamique : la section "Préalables" n'apparaît que si CONDITIONAL.
    from itertools import count as _count
    _counter = _count(1)
    def _next_num() -> str:
        return str(next(_counter))

    # ── RECOMMANDATION & JUSTIFICATION ────────────────────────────────────────
    h["section_title"]("Recommandation & Justification", num=_next_num())
    if scoring.justification:
        jp = doc.add_paragraph()
        jp.paragraph_format.space_after = Pt(8)
        jp.paragraph_format.alignment = WD.JUSTIFY
        jr = jp.add_run(scoring.justification)
        jr.font.size = Pt(10); jr.italic = True; jr.font.color.rgb = h["C_TEXT"]

    # ── GRILLE D'ÉVALUATION DÉTAILLÉE ─────────────────────────────────────────
    if scoring.criteria_breakdown:
        h["section_title"]("Grille d'évaluation détaillée", num=_next_num(), color=h["C_NAVY"])
        total_est = sum(c.estimated_score for c in scoring.criteria_breakdown)
        total_max = sum(c.max_points for c in scoring.criteria_breakdown)
        grid_t = doc.add_table(rows=1, cols=5)
        grid_t.style = "Table Grid"
        for cell, lbl in zip(grid_t.rows[0].cells, ["Critère", "Max", "Estimé", "Risque", "Justification"]):
            cell.text = lbl
            h["cell_bg"](cell, "0F295A")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(10)
        # Code couleur par niveau de risque du critère
        risk_color = {"FAIBLE": "DCFCE7", "MODÉRÉ": "FEF9C3", "ÉLEVÉ": "FFEDD5", "CRITIQUE": "FEE2E2"}
        for i, c in enumerate(scoring.criteria_breakdown):
            row = grid_t.add_row().cells
            # Critère inféré (non trouvé verbatim) → marqué « (estimé) »
            row[0].text = c.label + (" (estimé)" if c.is_inferred else "")
            row[1].text = str(c.max_points)
            row[2].text = str(c.estimated_score)
            row[3].text = c.risk_level
            row[4].text = c.rationale
            zebra = "FFFFFF" if i % 2 == 0 else "F8FAFC"
            h["cell_bg"](row[0], zebra); h["cell_bg"](row[1], zebra); h["cell_bg"](row[2], zebra)
            h["cell_bg"](row[3], risk_color.get(c.risk_level, "F1F5F9")); h["cell_bg"](row[4], zebra)
            for cell in row:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)
            # Label du critère en italique si inféré (signale qu'il n'est pas dans l'AO)
            if c.is_inferred and row[0].paragraphs[0].runs:
                row[0].paragraphs[0].runs[0].italic = True
            row[0].width = Cm(6.5); row[1].width = Cm(1.3); row[2].width = Cm(1.5)
            row[3].width = Cm(2.2); row[4].width = Cm(5)
        # Ligne de total
        tot = grid_t.add_row().cells
        tot[0].text = "TOTAL"
        tot[1].text = str(total_max)
        tot[2].text = str(total_est)
        tot[3].text = f"{scoring.score}/100"
        tot[4].text = ""
        for cell in tot:
            h["cell_bg"](cell, "E2E8F0")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.size = Pt(9); run.font.color.rgb = h["C_NAVY"]
        doc.add_paragraph("")

    # ── RISQUES & MITIGATION ──────────────────────────────────────────────────
    if scoring.risks:
        h["section_title"]("Risques & mitigation", num=_next_num(), color=h["C_RED"])
        risk_t = doc.add_table(rows=1, cols=4)
        risk_t.style = "Table Grid"
        for cell, lbl in zip(risk_t.rows[0].cells, ["Risque", "Criticité", "Pourquoi", "Mitigation"]):
            cell.text = lbl
            h["cell_bg"](cell, "B91C1C")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(10)
        crit_color = {"MODÉRÉ": "FEF9C3", "ÉLEVÉ": "FFEDD5", "CRITIQUE": "FEE2E2", "BLOQUANT": "FECACA"}
        for i, r in enumerate(scoring.risks):
            row = risk_t.add_row().cells
            row[0].text = r.label
            row[1].text = r.criticite
            row[2].text = r.pourquoi
            row[3].text = r.mitigation or "— (à définir)"
            zebra = "FFFFFF" if i % 2 == 0 else "F8FAFC"
            h["cell_bg"](row[0], zebra); h["cell_bg"](row[1], crit_color.get(r.criticite, "F1F5F9"))
            h["cell_bg"](row[2], zebra); h["cell_bg"](row[3], "F0FDF4")
            for cell in row:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)
            if row[1].paragraphs[0].runs:
                row[1].paragraphs[0].runs[0].bold = True
            row[0].width = Cm(4.5); row[1].width = Cm(2); row[2].width = Cm(4.5); row[3].width = Cm(5.5)
        doc.add_paragraph("")

    # ── PRÉALABLES CONDITIONNELS (seulement si CONDITIONAL) ───────────────────
    if scoring.recommendation.value == "CONDITIONAL":
        h["section_title"]("Préalables conditionnels", num=_next_num(), color=h["C_AMBER"])
        if scoring.preconditions:
            intro = doc.add_paragraph()
            intro.paragraph_format.space_after = Pt(6)
            ir = intro.add_run("Conditions à lever pour faire passer la recommandation de CONDITIONNEL à GO :")
            ir.font.size = Pt(10); ir.italic = True; ir.font.color.rgb = h["C_TEXT"]
            for p in scoring.preconditions:
                line = doc.add_paragraph()
                line.paragraph_format.left_indent = Cm(0.6)
                line.paragraph_format.space_after = Pt(3)
                box = line.add_run("☐  ")
                box.font.size = Pt(11)
                box.font.color.rgb = h["C_RED"] if p.blocking else h["C_AMBER"]
                meta = f"[{p.type}]"
                if p.deadline:
                    meta += f" {p.deadline}"
                if p.responsable:
                    meta += f" · {p.responsable}"
                lbl_r = line.add_run(p.label + " ")
                lbl_r.font.size = Pt(10); lbl_r.font.color.rgb = h["C_TEXT"]
                if p.blocking:
                    lbl_r.bold = True
                meta_r = line.add_run(meta)
                meta_r.font.size = Pt(8); meta_r.font.color.rgb = h["C_GRAY"]; meta_r.italic = True
        else:
            # Garde-fou : CONDITIONAL sans préalables (dégradation gracieuse côté pipeline)
            warn = doc.add_paragraph()
            warn.paragraph_format.space_after = Pt(6)
            wr = warn.add_run(
                "⚠ Recommandation conditionnelle sans préalables explicites — "
                "à compléter manuellement avant décision."
            )
            wr.font.size = Pt(10); wr.italic = True; wr.font.color.rgb = h["C_AMBER"]
        doc.add_paragraph("")

    # ── FORCES & ATOUTS ───────────────────────────────────────────────────────
    if scoring.strengths:
        h["section_title"]("Forces & Atouts de Neurones", num=_next_num(), color=h["C_GREEN"])
        for s in scoring.strengths:
            h["bullet"](s, h["C_GREEN"])
        doc.add_paragraph("")

    # ── ANALYSE DES ÉCARTS ────────────────────────────────────────────────────
    if scoring.gaps_analysis:
        h["section_title"]("Analyse des écarts", num=_next_num(), color=h["C_AMBER"])
        h["render_rich_text"](scoring.gaps_analysis)
        doc.add_paragraph("")

    # ── DOCUMENTS GED PERTINENTS ──────────────────────────────────────────────
    all_docs = list(scoring.team_matches or []) + list(scoring.similar_projects or []) + list(scoring.matched_documents or [])
    seen_keys: set = set()
    unique_docs = []
    for d in all_docs:
        key = d.doc_id or d.filename
        if key not in seen_keys:
            seen_keys.add(key)
            unique_docs.append(d)
    unique_docs.sort(key=lambda d: d.relevance_score, reverse=True)
    if unique_docs:
        h["section_title"]("Documents GED pertinents", num=_next_num(), color=h["C_NAVY"])
        ged_t = doc.add_table(rows=1, cols=3)
        ged_t.style = "Table Grid"
        for cell, lbl in zip(ged_t.rows[0].cells, ["Document", "Type", "Pertinence"]):
            cell.text = lbl
            h["cell_bg"](cell, "1E40AF")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(10)
        for i, d in enumerate(unique_docs[:10]):
            row = ged_t.add_row().cells
            row[0].text = d.filename
            row[1].text = str(d.doc_type)
            score_pct = int(d.relevance_score * 100) if d.relevance_score <= 1.0 else int(d.relevance_score)
            row[2].text = f"{score_pct}%"
            bg = "F0FDF4" if i % 2 == 0 else "FFFFFF"
            for cell in row:
                h["cell_bg"](cell, bg)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(10)
        doc.add_paragraph("")

    h["footer_line"]()
    buffer = BytesIO()
    doc.save(buffer)
    safe_name = re.sub(r"\.(pdf|docx)$", "", scoring.ao_filename, flags=re.IGNORECASE).replace(" ", "-")
    return Response(
        content=buffer.getvalue(), media_type=_DOCX_MIME,
        headers=_attachment_headers(f"Scoring-AO_{safe_name}.docx"),
    )


@router.post("/export-strategy")
async def export_strategy(body: StrategyExportRequest):
    """Exporte la stratégie de réponse en document Word professionnel (7 sections)."""
    from docx import Document as DocxDocument

    scoring = _from_schema(body.scoring_result)
    client_name = body.client_name or "Non précisé"
    dec_labels = {"GO": "GO — Répondre à l'AO", "CONDITIONAL": "CONDITIONNEL — Sous réserve", "NO_BID": "NO-BID — Ne pas répondre"}
    decision_label = dec_labels.get(body.decision, body.decision)

    doc = DocxDocument()
    h = _docx_helpers(doc)
    Pt = h["Pt"]; Cm = h["Cm"]
    WD = h["WD_ALIGN_PARAGRAPH"]

    rec_map = {
        "GO": ("GO — Répondre à l'AO", h["C_GREEN"]),
        "CONDITIONAL": ("CONDITIONNEL — Sous réserve", h["C_AMBER"]),
        "NO_BID": ("NO-BID — Ne pas répondre", h["C_RED"]),
    }
    rec_label, rec_color = rec_map.get(scoring.recommendation.value, ("—", h["C_AMBER"]))
    dec_color = h["C_GREEN"] if body.decision == "GO" else (h["C_AMBER"] if body.decision == "CONDITIONAL" else h["C_RED"])

    # ── PAGE DE COUVERTURE ────────────────────────────────────────────────────
    t1 = doc.add_paragraph()
    t1.alignment = WD.CENTER
    t1.paragraph_format.space_before = Pt(30)
    t1.paragraph_format.space_after = Pt(2)
    t1r = t1.add_run("STRATÉGIE DE RÉPONSE")
    t1r.bold = True; t1r.font.size = Pt(28); t1r.font.color.rgb = h["C_NAVY"]

    t2 = doc.add_paragraph()
    t2.alignment = WD.CENTER
    t2.paragraph_format.space_after = Pt(4)
    t2r = t2.add_run("Appel d'offres — " + scoring.ao_filename)
    t2r.italic = True; t2r.font.size = Pt(12); t2r.font.color.rgb = h["C_BLUE"]

    sep = doc.add_paragraph()
    sep.alignment = WD.CENTER
    sep.paragraph_format.space_after = Pt(20)
    sep.add_run("─" * 45).font.color.rgb = h["C_BLUE"]

    h["cover_table"]([
        ("Client / Commanditaire", client_name),
        ("Score de matching GED", f"{scoring.score} / 100"),
        ("Recommandation IA", rec_label),
        ("Décision commerciale", decision_label),
    ])

    dec_p = doc.add_paragraph()
    dec_p.alignment = WD.CENTER
    dec_p.paragraph_format.space_before = Pt(20)
    dec_p.paragraph_format.space_after = Pt(30)
    dec_r = dec_p.add_run(f"DÉCISION : {decision_label.upper()}")
    dec_r.bold = True; dec_r.font.size = Pt(14); dec_r.font.color.rgb = dec_color

    doc.add_page_break()

    bid = body.bid_strategy
    from itertools import count as _count
    _counter = _count(1)
    def _next_num() -> str:
        return str(next(_counter))

    # ── VALIDATION PARTENAIRE / PRÉALABLES (étape 0, si applicable) ────────────
    if bid.partner_validation:
        title = "Validation partenaire (étape 0)" if bid.partner else "Préalables à lever (étape 0)"
        h["section_title"](title, num=_next_num(), color=h["C_AMBER"])
        if bid.partner:
            pp = doc.add_paragraph()
            pp.paragraph_format.space_after = Pt(6)
            pr = pp.add_run(f"Partenaire de groupement : {bid.partner.name} "
                            f"({bid.partner.role}, {bid.partner.type})")
            pr.font.size = Pt(10); pr.bold = True; pr.font.color.rgb = h["C_NAVY"]
        pv_t = doc.add_table(rows=1, cols=5)
        pv_t.style = "Table Grid"
        for cell, lbl in zip(pv_t.rows[0].cells, ["Critère", "Type", "Pièces requises", "Bloquant", "Statut"]):
            cell.text = lbl
            h["cell_bg"](cell, "92400E")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(10)
        for i, pv in enumerate(bid.partner_validation):
            row = pv_t.add_row().cells
            row[0].text = pv.label
            row[1].text = pv.type
            row[2].text = "\n".join(pv.pieces_requises) if pv.pieces_requises else "—"
            row[3].text = "OUI" if pv.blocking else "non"
            row[4].text = pv.status
            zebra = "FFFFFF" if i % 2 == 0 else "F8FAFC"
            # Code couleur : bloquant non levé = rouge clair
            crit_bg = "FECACA" if pv.blocking else "FEF9C3"
            h["cell_bg"](row[0], zebra); h["cell_bg"](row[1], zebra)
            h["cell_bg"](row[2], zebra); h["cell_bg"](row[3], crit_bg); h["cell_bg"](row[4], zebra)
            for cell in row:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)
            if row[3].paragraphs[0].runs:
                row[3].paragraphs[0].runs[0].bold = True
            row[0].width = Cm(4.5); row[1].width = Cm(2.2); row[2].width = Cm(5)
            row[3].width = Cm(1.8); row[4].width = Cm(2.5)
        doc.add_paragraph("")

    # ── CONTEXTE DE L'AO ──────────────────────────────────────────────────────
    h["section_title"]("Contexte de l'appel d'offres", num=_next_num())
    if scoring.summary:
        h["body_para"](scoring.summary)
    if scoring.key_elements:
        doc.add_paragraph("")
        kp_t = doc.add_table(rows=1, cols=2)
        kp_t.style = "Table Grid"
        for cell, lbl in zip(kp_t.rows[0].cells, ["Critère", "Valeur"]):
            cell.text = lbl
            h["cell_bg"](cell, "0F295A")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(10)
        for i, el in enumerate(scoring.key_elements):
            row = kp_t.add_row().cells
            row[0].text = el.category
            row[1].text = el.value
            h["cell_bg"](row[0], "EFF6FF" if i % 2 == 0 else "FFFFFF")
            for cell in row:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(10)
            if row[0].paragraphs[0].runs:
                row[0].paragraphs[0].runs[0].bold = True
            row[0].width = Cm(5.5); row[1].width = Cm(10.5)
        doc.add_paragraph("")

    # ── STRATÉGIE DE RÉPONSE ──────────────────────────────────────────────────
    h["section_title"]("Stratégie de réponse", num=_next_num())
    if bid.strategy_text.strip():
        h["render_rich_text"](bid.strategy_text)
    else:
        h["body_para"]("Stratégie non définie.")
    doc.add_paragraph("")

    # ── PLAN DE RÉPONSE EN PHASES ─────────────────────────────────────────────
    if bid.phases:
        h["section_title"]("Plan de réponse en phases", num=_next_num(), color=h["C_NAVY"])
        for ph in bid.phases:
            # En-tête de phase
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(8)
            hp.paragraph_format.space_after = Pt(3)
            day_range = f" ({ph.start_day}→{ph.end_day})" if ph.start_day else ""
            hr = hp.add_run(f"{ph.id} · {ph.name}{day_range}")
            hr.bold = True; hr.font.size = Pt(10); hr.font.color.rgb = h["C_BLUE"]
            if ph.is_blocking_next:
                br = hp.add_run("   ⛔ bloque la phase suivante tant que non terminée")
                br.font.size = Pt(8); br.italic = True; br.font.color.rgb = h["C_RED"]
            if not ph.actions:
                h["body_para"]("Actions à préciser.")
                continue
            ph_t = doc.add_table(rows=1, cols=4)
            ph_t.style = "Table Grid"
            for cell, lbl in zip(ph_t.rows[0].cells, ["Jour", "Action", "Responsable", "Livrable"]):
                cell.text = lbl
                h["cell_bg"](cell, "1E40AF")
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(9)
            for i, a in enumerate(ph.actions):
                row = ph_t.add_row().cells
                row[0].text = a.day_label
                row[1].text = a.action
                row[2].text = a.responsable
                row[3].text = a.deliverable or "—"
                zebra = "EFF6FF" if i % 2 == 0 else "FFFFFF"
                h["cell_bg"](row[0], "DBEAFE")
                h["cell_bg"](row[1], zebra); h["cell_bg"](row[2], zebra); h["cell_bg"](row[3], zebra)
                for cell in row:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(9)
                if row[0].paragraphs[0].runs:
                    row[0].paragraphs[0].runs[0].bold = True
                    row[0].paragraphs[0].runs[0].font.color.rgb = h["C_BLUE"]
                row[0].width = Cm(1.8); row[1].width = Cm(7.5); row[2].width = Cm(3.5); row[3].width = Cm(3.2)
        doc.add_paragraph("")

    # ── PLAN DE RÉPONSE DÉTAILLÉ ──────────────────────────────────────────────
    if bid.response_plan.strip():
        h["section_title"]("Plan de réponse détaillé", num=_next_num())
        h["render_rich_text"](bid.response_plan)
        doc.add_paragraph("")

    # ── CHECKLIST DES PIÈCES & ANNEXES ────────────────────────────────────────
    if bid.appendices:
        h["section_title"]("Checklist des pièces & annexes", num=_next_num(), color=h["C_NAVY"])
        ap_t = doc.add_table(rows=1, cols=6)
        ap_t.style = "Table Grid"
        for cell, lbl in zip(ap_t.rows[0].cells, ["Code", "Pièce", "Type", "Responsable", "Langue", "Statut"]):
            cell.text = lbl
            h["cell_bg"](cell, "0F295A")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(9)
        for i, ap in enumerate(bid.appendices):
            row = ap_t.add_row().cells
            row[0].text = ap.code
            row[1].text = ap.label + ("" if ap.obligatoire else "  (facultatif)")
            row[2].text = ap.type
            row[3].text = ap.responsable or "—"
            row[4].text = ap.langue or "—"
            row[5].text = ap.statut
            zebra = "EFF6FF" if i % 2 == 0 else "FFFFFF"
            for cell in row:
                h["cell_bg"](cell, zebra)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)
            if row[0].paragraphs[0].runs:
                row[0].paragraphs[0].runs[0].bold = True
            row[0].width = Cm(1.8); row[1].width = Cm(6); row[2].width = Cm(2.2)
            row[3].width = Cm(3); row[4].width = Cm(1.8); row[5].width = Cm(1.7)
        doc.add_paragraph("")

    # Numéro de version du plan
    ver_p = doc.add_paragraph()
    ver_p.alignment = WD.CENTER
    ver_p.paragraph_format.space_before = Pt(10)
    vr = ver_p.add_run(f"Version {bid.version} du plan de réponse"
                       + (f" · généré le {bid.generated_at}" if bid.generated_at else ""))
    vr.font.size = Pt(8); vr.italic = True; vr.font.color.rgb = h["C_GRAY"]

    h["footer_line"]()
    buffer = BytesIO()
    doc.save(buffer)
    safe_name = re.sub(r"\.(pdf|docx)$", "", scoring.ao_filename, flags=re.IGNORECASE).replace(" ", "-")
    return Response(
        content=buffer.getvalue(), media_type=_DOCX_MIME,
        headers=_attachment_headers(f"Strategie-Reponse_{safe_name}.docx"),
    )


@router.post("/export-checklist")
async def export_checklist(body: ChecklistExportRequest):
    """Exporte la checklist dossier en document Word professionnel."""
    from docx import Document as DocxDocument

    doc = DocxDocument()
    h = _docx_helpers(doc)
    Pt = h["Pt"]; Cm = h["Cm"]
    WD = h["WD_ALIGN_PARAGRAPH"]

    # ── Titre ────────────────────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.alignment = WD.CENTER
    title_p.paragraph_format.space_before = Pt(20)
    title_p.paragraph_format.space_after = Pt(2)
    tr = title_p.add_run("CHECKLIST DOSSIER DE RÉPONSE")
    tr.bold = True; tr.font.size = Pt(22); tr.font.color.rgb = h["C_NAVY"]

    sep = doc.add_paragraph()
    sep.alignment = WD.CENTER
    sep.paragraph_format.space_after = Pt(16)
    sep.add_run("─" * 45).font.color.rgb = h["C_BLUE"]

    h["cover_table"]([
        ("Appel d'offres", body.ao_filename),
        ("Client / Commanditaire", body.client_name or "Non précisé"),
        ("Date de soumission prévue", body.submission_date or "Non précisée"),
        ("Date d'export", datetime.now().strftime("%d %B %Y")),
    ])

    # Pièces réelles de l'AO ; sinon checklist générique inférée
    appendices = list(body.appendices)
    inferred = not appendices
    if inferred:
        appendices = _GENERIC_CHECKLIST

    # Score de complétude : pièces obligatoires au statut OK / total obligatoire
    obligatoires = [a for a in appendices if a.obligatoire]
    ok_oblig = [a for a in obligatoires if a.statut == "OK"]
    ok_total = sum(1 for a in appendices if a.statut == "OK")
    score_p = doc.add_paragraph()
    score_p.alignment = WD.CENTER
    score_p.paragraph_format.space_before = Pt(14)
    score_p.paragraph_format.space_after = Pt(6)
    n_oblig = len(obligatoires)
    score_color = h["C_GREEN"] if n_oblig and len(ok_oblig) == n_oblig else (
        h["C_AMBER"] if n_oblig and len(ok_oblig) >= n_oblig * 0.6 else h["C_RED"]
    )
    sc_r = score_p.add_run(
        f"Complétude : {len(ok_oblig)}/{n_oblig} pièces obligatoires fournies"
        f"  ·  {ok_total}/{len(appendices)} pièces totales"
    )
    sc_r.bold = True; sc_r.font.size = Pt(11); sc_r.font.color.rgb = score_color

    if inferred:
        warn_p = doc.add_paragraph()
        warn_p.alignment = WD.CENTER
        warn_p.paragraph_format.space_after = Pt(14)
        wr = warn_p.add_run(
            "⚠ Checklist générique (inférée) — l'AO ne listait pas de pièces explicites. "
            "À compléter et vérifier manuellement."
        )
        wr.font.size = Pt(9); wr.italic = True; wr.font.color.rgb = h["C_AMBER"]

    doc.add_page_break()

    # ── Sections par type de pièce ────────────────────────────────────────────────
    type_order = ["ADMIN", "TECHNIQUE", "FINANCIER", "RH"]
    type_label = {
        "ADMIN": "Pièces administratives",
        "TECHNIQUE": "Pièces techniques",
        "FINANCIER": "Pièces financières",
        "RH": "Pièces RH / équipe",
    }
    type_color = {
        "ADMIN": h["C_NAVY"], "TECHNIQUE": h["C_BLUE"],
        "FINANCIER": h["C_GREEN"], "RH": h["C_AMBER"],
    }
    # Symbole + couleur + fond par statut
    statut_symbol = {"OK": "✓", "NOK": "✗", "EN_COURS": "◐", "PENDING": "☐"}
    statut_fg = {"OK": h["C_GREEN"], "NOK": h["C_RED"], "EN_COURS": h["C_AMBER"], "PENDING": h["C_RED"]}
    statut_bg = {"OK": "D1FAE5", "NOK": "FEE2E2", "EN_COURS": "FEF9C3"}

    from itertools import count as _count
    _counter = _count(1)

    for ptype in type_order:
        items_type = [a for a in appendices if a.type == ptype]
        if not items_type:
            continue
        h["section_title"](type_label[ptype], color=type_color[ptype], num=str(next(_counter)))

        tbl = doc.add_table(rows=1, cols=4)
        tbl.style = "Table Grid"
        for cell, lbl in zip(tbl.rows[0].cells, ["Statut", "Code", "Pièce du dossier", "Observations"]):
            cell.text = lbl
            h["cell_bg"](cell, "0F295A")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True; run.font.color.rgb = h["C_WHITE"]; run.font.size = Pt(9)

        for i, item in enumerate(items_type):
            row = tbl.add_row().cells
            row[0].text = statut_symbol.get(item.statut, "☐")
            row[1].text = item.code or "—"
            row[2].text = item.label + ("" if item.obligatoire else "  (facultatif)")
            row[3].text = item.note or (item.source_section or "")

            default_bg = "FFFFFF" if i % 2 == 0 else "F8FAFC"
            bg = statut_bg.get(item.statut, default_bg)
            for cell in row:
                h["cell_bg"](cell, bg)

            for para in row[0].paragraphs:
                para.paragraph_format.alignment = WD.CENTER
                for run in para.runs:
                    run.font.size = Pt(12); run.bold = True
                    run.font.color.rgb = statut_fg.get(item.statut, h["C_RED"])
            for para in row[1].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9); run.bold = True
            for para in row[2].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)
                    if item.obligatoire:
                        run.bold = True
            for para in row[3].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9); run.font.color.rgb = h["C_GRAY"]

            row[0].width = Cm(1.5); row[1].width = Cm(2); row[2].width = Cm(9); row[3].width = Cm(4)

        doc.add_paragraph("")

    h["footer_line"]()
    buffer = BytesIO()
    doc.save(buffer)
    safe_name = re.sub(r"\.(pdf|docx)$", "", body.ao_filename, flags=re.IGNORECASE).replace(" ", "-")
    return Response(
        content=buffer.getvalue(), media_type=_DOCX_MIME,
        headers=_attachment_headers(f"Checklist_{safe_name}.docx"),
    )


@router.post("/match-team", response_model=TeamMatchResponse)
async def match_team(body: TeamMatchRequest, request: Request):
    """Cherche dans la GED les CVs correspondant aux profils demandés par l'AO."""
    use_case = _get_use_case(request)
    return await use_case.match_team(body)


def _to_schema(result: ScoringResult) -> ScoringResultSchema:
    return ScoringResultSchema(
        ao_filename=result.ao_filename,
        summary=result.summary,
        key_elements=[KeyElementSchema(category=e.category, value=e.value) for e in result.key_elements],
        matched_documents=[
            MatchedDocumentSchema(
                doc_id=m.doc_id,
                filename=m.filename,
                doc_type=m.doc_type,
                relevance_score=m.relevance_score,
                excerpt=m.excerpt,
            )
            for m in result.matched_documents
        ],
        gaps_analysis=result.gaps_analysis,
        strengths=result.strengths,
        risks=[RiskSchema(**r.__dict__) for r in result.risks],
        score=result.score,
        score_basis=result.score_basis,
        recommendation=BidRecommendationSchema(result.recommendation.value),
        justification=result.justification,
        criteres_selection=result.criteres_selection,
        besoins=result.besoins,
        prerequis=result.prerequis,
        ressources_demandees=result.ressources_demandees,
        points_vigilance=result.points_vigilance,
        date_remise=result.date_remise,
        team_matches=[
            MatchedDocumentSchema(
                doc_id=m.doc_id, filename=m.filename, doc_type=m.doc_type,
                relevance_score=m.relevance_score, excerpt=m.excerpt,
            )
            for m in result.team_matches
        ],
        similar_projects=[
            MatchedDocumentSchema(
                doc_id=m.doc_id, filename=m.filename, doc_type=m.doc_type,
                relevance_score=m.relevance_score, excerpt=m.excerpt,
            )
            for m in result.similar_projects
        ],
        market_identity=MarketIdentitySchema(**result.market_identity.__dict__),
        calendar=[CalendarEventSchema(**c.__dict__) for c in result.calendar],
        evaluation_modalities=EvaluationModalitiesSchema(**result.evaluation_modalities.__dict__),
        criteria_breakdown=[ScoringCriterionSchema(**c.__dict__) for c in result.criteria_breakdown],
        preconditions=[PreconditionSchema(**p.__dict__) for p in result.preconditions],
        preconditions_incomplete=result.preconditions_incomplete,
        appendices=[AppendixSchema(**a.__dict__) for a in result.appendices],
        appendices_incomplete=result.appendices_incomplete,
        profils_demandes=[RequiredProfileSchema(**p.__dict__) for p in result.profils_demandes],
        seuils_eligibilite=[EligibilityThresholdSchema(**s.__dict__) for s in result.seuils_eligibilite],
        donnees_financieres=FinancialDataSchema(**result.donnees_financieres.__dict__),
        client_context=ClientContextSchema(**result.client_context.__dict__),
        capability_matches=[
            CapabilityMatchSchema(
                theme=m.theme, confidence=m.confidence, is_critical=m.is_critical,
                won_count=m.won_count, clients=m.clients,
                deals=[CapabilityDealSchema(**d.__dict__) for d in m.deals],
            )
            for m in result.capability_matches
        ],
        capability_gaps=result.capability_gaps,
    )


def _from_schema(schema: ScoringResultSchema) -> ScoringResult:
    from core.domain.offer import (
        KeyElement, MatchedDocument, BidRecommendation,
        MarketIdentity, CalendarEvent, EvaluationModalities,
        ScoringCriterion, Risk, Precondition, Appendix,
        RequiredProfile, EligibilityThreshold, FinancialData, ExtractedItem,
        ClientContext, CapabilityMatch, CapabilityDeal,
    )
    return ScoringResult(
        ao_filename=schema.ao_filename,
        summary=schema.summary,
        key_elements=[KeyElement(category=e.category, value=e.value) for e in schema.key_elements],
        matched_documents=[
            MatchedDocument(
                doc_id=m.doc_id, filename=m.filename, doc_type=m.doc_type,
                relevance_score=m.relevance_score, excerpt=m.excerpt,
            )
            for m in schema.matched_documents
        ],
        gaps_analysis=schema.gaps_analysis,
        strengths=schema.strengths,
        risks=[Risk(**r.model_dump()) for r in schema.risks],
        score=schema.score,
        score_basis=schema.score_basis,
        recommendation=BidRecommendation(schema.recommendation.value),
        justification=schema.justification,
        criteres_selection=[ExtractedItem.coerce(i) for i in schema.criteres_selection],
        besoins=[ExtractedItem.coerce(i) for i in schema.besoins],
        prerequis=[ExtractedItem.coerce(i) for i in schema.prerequis],
        ressources_demandees=[ExtractedItem.coerce(i) for i in schema.ressources_demandees],
        points_vigilance=[ExtractedItem.coerce(i) for i in schema.points_vigilance],
        date_remise=schema.date_remise,
        team_matches=[
            MatchedDocument(
                doc_id=m.doc_id, filename=m.filename, doc_type=m.doc_type,
                relevance_score=m.relevance_score, excerpt=m.excerpt,
            )
            for m in schema.team_matches
        ],
        similar_projects=[
            MatchedDocument(
                doc_id=m.doc_id, filename=m.filename, doc_type=m.doc_type,
                relevance_score=m.relevance_score, excerpt=m.excerpt,
            )
            for m in schema.similar_projects
        ],
        market_identity=MarketIdentity(**schema.market_identity.model_dump()),
        calendar=[CalendarEvent(**c.model_dump()) for c in schema.calendar],
        evaluation_modalities=EvaluationModalities(**schema.evaluation_modalities.model_dump()),
        criteria_breakdown=[ScoringCriterion(**c.model_dump()) for c in schema.criteria_breakdown],
        preconditions=[Precondition(**p.model_dump()) for p in schema.preconditions],
        preconditions_incomplete=schema.preconditions_incomplete,
        appendices=[Appendix(**a.model_dump()) for a in schema.appendices],
        appendices_incomplete=schema.appendices_incomplete,
        profils_demandes=[RequiredProfile(**p.model_dump()) for p in schema.profils_demandes],
        seuils_eligibilite=[EligibilityThreshold(**s.model_dump()) for s in schema.seuils_eligibilite],
        donnees_financieres=FinancialData(**schema.donnees_financieres.model_dump()),
        client_context=ClientContext(**schema.client_context.model_dump()),
        capability_matches=[
            CapabilityMatch(
                theme=m.theme, confidence=m.confidence, is_critical=m.is_critical,
                won_count=m.won_count, clients=m.clients,
                deals=[CapabilityDeal(**d.model_dump()) for d in m.deals],
            )
            for m in schema.capability_matches
        ],
        capability_gaps=schema.capability_gaps,
    )
