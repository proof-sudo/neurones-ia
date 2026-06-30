"""Tests de l'export Excel de la matrice de conformité (C4).

On rend un .xlsx puis on le RÉOUVRE avec openpyxl pour vérifier la structure
(feuilles, en-tête, nombre de lignes = nb d'exigences, en-tête figé). Aucun LLM."""
from io import BytesIO

from openpyxl import load_workbook

from core.domain.offer import (
    ScoringResult, ExtractedItem, EligibilityThreshold, Appendix, BidRecommendation,
)
from modules.uc10_presales.requirements_builder import build_matrice
from modules.uc10_presales.matrix_export import render_matrix_xlsx


def _scoring() -> ScoringResult:
    return ScoringResult(
        ao_filename="AO-Demo.pdf", summary="", key_elements=[], matched_documents=[],
        gaps_analysis="", strengths=[], risks=[], score=60,
        recommendation=BidRecommendation.GO, justification="",
        besoins=[ExtractedItem(texte="Virtualisation de 50 serveurs", source_section="§4.1")],
        criteres_selection=[ExtractedItem(texte="3 références bancaires", source_section="§2")],
        seuils_eligibilite=[EligibilityThreshold(libelle="CA min", valeur="500M", unite="FCFA", blocking=True)],
        appendices=[Appendix(code="5A", label="Déclaration de conformité")],
    )


def test_render_produces_valid_xlsx_with_expected_structure():
    matrice = build_matrice(_scoring(), generated_at="2026-06-22T10:00:00")
    data = render_matrix_xlsx(matrice)
    assert data[:2] == b"PK"  # signature ZIP (xlsx = zip)

    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Matrice de conformité", "Synthèse"]

    ws = wb["Matrice de conformité"]
    # en-tête attendu
    assert ws.cell(row=1, column=1).value == "Réf."
    assert ws.cell(row=1, column=3).value == "Exigence"
    # une ligne de données par exigence (4 ici)
    assert ws.max_row == matrice.total + 1 == 5
    assert ws.freeze_panes == "A2"
    # le texte verbatim et la source sont présents
    textes = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
    assert "Virtualisation de 50 serveurs" in textes
    sources = [ws.cell(row=r, column=4).value for r in range(2, ws.max_row + 1)]
    assert "§4.1" in sources


def test_summary_sheet_reports_exhaustivity():
    matrice = build_matrice(_scoring())
    wb = load_workbook(BytesIO(render_matrix_xlsx(matrice)))
    ws = wb["Synthèse"]
    flat = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Total exigences" in flat
    assert matrice.total in flat
