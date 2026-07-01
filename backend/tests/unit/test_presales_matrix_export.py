"""Export Excel de la matrice : colonnes IA suggéré + confirmation humaine présentes."""
from io import BytesIO

from openpyxl import load_workbook

from core.domain.requirements import Exigence, MatriceConformite, STATUT_CONFORME, TYPE_CRITERE
from modules.uc10_presales.matrix_export import render_matrix_xlsx


def _matrice():
    return MatriceConformite(ao_filename="ao.pdf", exigences=[
        Exigence(id="EX-0001", texte="Trois références similaires", type=TYPE_CRITERE,
                 statut_suggere=STATUT_CONFORME, justification_ia="3 PV en GED", confiance_ia=0.8),
    ], expected_total=1)


def test_export_contient_colonnes_ia_et_confirmation():
    data = render_matrix_xlsx(_matrice())
    wb = load_workbook(BytesIO(data))
    ws = wb["Matrice de conformité"]
    headers = [c.value for c in ws[1]]
    assert "Statut suggéré (IA)" in headers
    assert "Justification IA" in headers
    assert "Statut validé (humain)" in headers
    assert "Confirmé par" in headers


def test_export_remplit_le_statut_suggere_et_laisse_le_valide_vide():
    data = render_matrix_xlsx(_matrice())
    ws = load_workbook(BytesIO(data))["Matrice de conformité"]
    headers = [c.value for c in ws[1]]
    row = {headers[i]: ws.cell(row=2, column=i + 1).value for i in range(len(headers))}
    assert row["Statut suggéré (IA)"] == "Conforme"
    assert row["Justification IA"] == "3 PV en GED"
    # Non coché par l'humain → le statut validé reste vide (on n'affirme rien sans cochage).
    assert (row["Statut validé (humain)"] or "") == ""
