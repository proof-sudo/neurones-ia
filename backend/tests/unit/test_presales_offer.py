"""Offre sans template (UC10) : contexte enrichi + parse tolérant du plan de document
+ repli déterministe + construction .docx de bout en bout."""
from io import BytesIO

from docx import Document as DocxDocument

from core.domain.offer import ScoringResult, ExtractedItem, Risk, BidRecommendation
from modules.uc10_presales.offer_generator import (
    _offer_user_context, _parse_document, _fallback_document, _sanitize_block,
)
from modules.uc10_presales.offer_docx_builder import build_offer_docx, _named_team_from_cvs


def _scoring(**kw) -> ScoringResult:
    base = dict(
        ao_filename="ao.pdf", summary="Refonte du SI bancaire", key_elements=[],
        matched_documents=[], gaps_analysis="Pas de référence BI récente", strengths=[],
        risks=[], score=62, recommendation=BidRecommendation.GO, justification="",
    )
    base.update(kw)
    return ScoringResult(**base)


# ── Parse du plan de document ─────────────────────────────────────────────────

def test_parse_document_valide():
    raw = (
        '{"cover": {"titre": "Déploiement Odoo", "sous_titre": "Offre technique"},'
        ' "blocks": [{"type": "heading", "level": 1, "text": "Besoin"},'
        ' {"type": "paragraph", "text": "Contexte."},'
        ' {"type": "bullets", "items": ["a", "b"]}]}'
    )
    doc = _parse_document(raw)
    assert doc["cover"]["titre"] == "Déploiement Odoo"
    assert [b["type"] for b in doc["blocks"]] == ["heading", "paragraph", "bullets"]


def test_parse_document_tronque_recupere_les_blocs():
    # JSON coupé en plein milieu d'un bloc : les blocs déjà fermés doivent survivre.
    raw = (
        '{"cover": {"titre": "Plateforme bancaire"},'
        ' "blocks": [{"type": "heading", "level": 1, "text": "Compréhension"},'
        ' {"type": "paragraph", "text": "BGFI, secteur bancaire."},'
        ' {"type": "table", "titre": "Stack", "headers": ["Composant'  # ← tronqué ici
    )
    doc = _parse_document(raw)
    assert doc["cover"]["titre"].startswith("Plateforme")
    assert len(doc["blocks"]) == 2  # heading + paragraph récupérés, table tronquée ignorée
    assert doc["blocks"][0]["type"] == "heading"


def test_parse_document_garbage_renvoie_vide():
    assert _parse_document("ceci n'est pas du JSON du tout") == {}


def test_sanitize_bloc_type_inconnu_ignore():
    assert _sanitize_block({"type": "video", "src": "x"}) is None
    assert _sanitize_block({"type": "heading", "text": ""}) is None  # heading vide rejeté
    assert _sanitize_block({"type": "heading", "level": 9, "text": "T"})["level"] == 3


# ── Repli déterministe ────────────────────────────────────────────────────────

def test_fallback_document_non_vide():
    sc = _scoring(besoins=[ExtractedItem("Migration comptable")], strengths=["Expertise Odoo"])
    doc = _fallback_document(sc, "BGFI")
    assert doc["cover"]["titre"]
    assert doc["blocks"]  # jamais vide
    texts = " ".join(b.get("text", "") + " ".join(b.get("items", [])) for b in doc["blocks"])
    assert "Migration comptable" in texts


# ── Contexte utilisateur ──────────────────────────────────────────────────────

def test_contexte_enrichi_contient_atouts_et_risques():
    sc = _scoring(
        strengths=["Expertise Odoo multi-sociétés confirmée"],
        risks=[Risk(label="Absence de référence BI", criticite="ÉLEVÉ", mitigation="Partenariat BI")],
        besoins=[ExtractedItem("Migration des données comptables")],
    )
    ctx = _offer_user_context(sc, "BGFI", "Digitalisation", "30 jours", "10M FCFA")
    assert "Expertise Odoo multi-sociétés confirmée" in ctx
    assert "ATOUTS" in ctx and "RISQUES" in ctx
    assert "Partenariat BI" in ctx
    assert "BGFI" in ctx


# ── Construction du .docx (sans template) ─────────────────────────────────────

def test_build_offer_docx_produit_un_docx_valide():
    cover = {"titre": "Déploiement d'une plateforme", "sous_titre": "Offre technique",
             "accroche": "Une solution sur mesure."}
    blocks = [
        {"type": "heading", "level": 1, "text": "Compréhension du besoin"},
        {"type": "paragraph", "text": "Le client souhaite moderniser son SI."},
        {"type": "bullets", "items": ["Point 1", "Point 2"]},
        {"type": "table", "titre": "Stack technique",
         "headers": ["Composant", "Version"], "rows": [["PostgreSQL", "15"], ["Nginx", "—"]]},
    ]
    data = build_offer_docx(cover=cover, blocks=blocks, client_name="BGFI Bank")
    assert isinstance(data, bytes) and len(data) > 0
    doc = DocxDocument(BytesIO(data))
    full = "\n".join(p.text for p in doc.paragraphs)
    # Titre couverture + bloc IA + blocs cannés (présentation/méthodologie/certifications) + équipe
    assert "Déploiement d'une plateforme" in full
    assert "Compréhension du besoin" in full
    assert "Présentation de Neurones Technologies" in full
    assert "Méthodologie et gestion de projet" in full
    assert "Équipe projet dédiée" in full
    # Le tableau IA + le tableau équipe sont présents
    assert len(doc.tables) >= 2


# ── Tableau équipe : priorité aux vraies données kb_cv ────────────────────────

def test_equipe_utilise_kb_cv_quand_disponible():
    kb_team = {"CV_Jean_Konan.pdf": ("Jean Konan", "Architecte Cloud")}
    people = _named_team_from_cvs(
        ["CV_Jean_Konan.pdf", "CV_inconnu.pdf"], kb_team=kb_team
    )
    assert people[0] == ("Jean Konan", "Architecte Cloud")
    # Le 2e CV n'a pas de correspondance kb_cv → repli sur le nom de fichier
    assert people[1][0] == "Inconnu"
    assert people[1][1] == "Consultant / Développeur Senior"


def test_equipe_repli_complet_sans_kb_team():
    people = _named_team_from_cvs(["CV_Awa_Diallo.pdf"], kb_team=None)
    assert people == [("Awa Diallo", "Chef de Projet")]


def test_build_offer_docx_accepte_kb_team():
    data = build_offer_docx(
        cover={"titre": "Test"}, blocks=[], client_name="Client",
        selected_cvs=["CV_Jean_Konan.pdf"],
        kb_team={"CV_Jean_Konan.pdf": ("Jean Konan", "Architecte Cloud")},
    )
    doc = DocxDocument(BytesIO(data))
    full = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "Jean Konan" in full
    assert "Architecte Cloud" in full
